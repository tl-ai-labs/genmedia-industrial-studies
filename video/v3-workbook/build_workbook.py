"""Build the video v3 scenario workbook (.xlsx) from src/*.yaml.

    voice/.venv/bin/python video/v3-workbook/build_workbook.py
    voice/.venv/bin/python video/v3-workbook/build_workbook.py --check      # validate only
    voice/.venv/bin/python video/v3-workbook/build_workbook.py --src DIR --out FILE

Same pattern as image/v3-workbook: YAML sources, a validating build, and a
workbook whose every count is derived from the sources.

    src/sources.yaml        video-only inputs (source clips and new images), model-neutral
    src/<fam>.yaml          one header per family: rubric weights and gate, targets, slot grid
    src/<fam>-<ind>.yaml    part files: `code` and `scenarios`, merged in file order

Video reuses the image bank's sources (image/v3-workbook/src/sources.yaml) as
first frames and references, so the image-to-video chain is built in. A base
scene's uses are counted across both lanes and capped at 5. Both banks are scored
with one rubric (image/v3-workbook/src/rubric.yaml); video uses the criteria
marked for video, including motion coherence and temporal consistency.

The bank is written one industry at a time. Rules that can only hold for a full
family or a full bank are reported as "pending" until then; the build stops on
any rule that is already broken.
"""
from __future__ import annotations

import argparse
import re
import sys
from collections import Counter, OrderedDict, defaultdict
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
SRC = HERE / "src"
IMAGE_SRC = ROOT / "image" / "v3-workbook" / "src"
OUT = ROOT / "video" / "docs" / "sheets" / "2026-10-08-video-v3-scenarios.xlsx"

# ---------------------------------------------------------------- the plan

INDUSTRIES = OrderedDict([
    ("FLM", "Film, OTT & micro-drama"),
    ("ADS", "Advertising & agencies"),
    ("ECM", "E-commerce, D2C & retail"),
    ("GAM", "Gaming"),
    ("RES", "Real estate, interiors & architecture"),
    ("MED", "Media, publishing & events"),
])

FAMILY_ORDER = ["VCOV", "VSTB", "VBTS", "VCHR", "VSPT", "VEDT", "VPHY", "VCAM", "VTXT"]
TASKS = OrderedDict([("t2v", "Text to video"), ("i2v", "Image to video"), ("edit", "Video edit")])

# The image family each video family is built on (for from_image links).
BUILT_ON = {"VCOV": "COV", "VSTB": "STB", "VBTS": "SHT", "VCHR": "CHR", "VSPT": "PRD",
            "VEDT": "EDT", "VPHY": "PHY", "VCAM": "CMP", "VTXT": "TYP"}

# Slots per family and industry: (scenarios, t2v, i2v, edit). Industries not listed get none.
PLAN = {
    "VCOV": {"FLM": (12, 3, 9, 0), "ADS": (5, 1, 4, 0), "ECM": (5, 1, 4, 0), "GAM": (6, 2, 4, 0), "RES": (9, 2, 7, 0), "MED": (3, 1, 2, 0)},
    "VSTB": {"FLM": (11, 0, 11, 0), "ADS": (6, 0, 6, 0), "ECM": (2, 0, 2, 0), "GAM": (6, 0, 6, 0)},
    "VBTS": {"FLM": (8, 0, 8, 0), "ADS": (7, 0, 7, 0), "ECM": (5, 0, 5, 0), "GAM": (5, 0, 5, 0)},
    "VCHR": {"FLM": (12, 3, 7, 2), "ADS": (4, 1, 2, 1), "ECM": (5, 1, 3, 1), "GAM": (9, 2, 6, 1)},
    "VSPT": {"ADS": (12, 2, 8, 2), "ECM": (14, 2, 9, 3), "GAM": (2, 1, 1, 0), "RES": (2, 0, 2, 0)},
    "VEDT": {"FLM": (7, 0, 0, 7), "ADS": (6, 0, 0, 6), "ECM": (9, 0, 0, 9), "GAM": (4, 0, 0, 4), "RES": (4, 0, 0, 4), "MED": (5, 0, 0, 5)},
    "VPHY": {"FLM": (5, 3, 1, 1), "ADS": (5, 3, 1, 1), "ECM": (3, 2, 0, 1), "GAM": (3, 2, 1, 0), "RES": (7, 5, 1, 1), "MED": (7, 5, 1, 1)},
    "VCAM": {"FLM": (4, 2, 1, 1), "ADS": (2, 1, 1, 0), "ECM": (3, 2, 0, 1), "GAM": (1, 1, 0, 0), "RES": (2, 1, 1, 0), "MED": (3, 1, 1, 1)},
    "VTXT": {"FLM": (1, 0, 0, 1), "ADS": (1, 1, 0, 0), "ECM": (2, 1, 0, 1), "MED": (6, 3, 0, 3)},
}
# Extreme scenarios per family (about 40%); the rest are hard.
EXTREME = {"VCOV": 16, "VSTB": 10, "VBTS": 10, "VCHR": 12, "VSPT": 12, "VEDT": 14, "VPHY": 12, "VCAM": 6, "VTXT": 4}
BANK_TOTAL = 240
MIN_SHARE = {"t2v": 0.20, "i2v": 0.45, "edit": 0.20}
MIN_INDIA = 0.30
MAX_SUBJECT_SHARE = 0.25
MIN_PER_FAMILY = {"industries": 4, "subjects": 6, "styles": 3, "lights": 4}
MAX_KIT_USES = 5          # per scene, image and video uses together
CHECKS_RANGE = (12, 25)
MEMBER_CHECKS_RANGE = (2, 5)
MIN_LEVERS = 3
MIN_LEVER_FAMILIES = 3
LEVER_FAMILY_EXEMPT = {"set": "only Coverage, Story beats and Character continuity produce sets"}
MIN_PROMPT_WORDS = 80
MIN_SOURCE_PROMPT_WORDS = 60
DURATION_RANGE = (3, 20)  # seconds per clip

SUBJECTS = OrderedDict([
    ("person", "Single person"), ("group", "Group (3+)"), ("child-elderly", "Child or elderly"),
    ("product", "Product"), ("food", "Food"), ("fashion", "Fashion"), ("vehicle", "Vehicle"),
    ("animal", "Animal"), ("interior", "Interior"), ("architecture", "Architecture"),
    ("landscape", "Landscape"), ("stylised", "Stylised character"),
])
STYLES = OrderedDict([
    ("photoreal", "Photoreal"), ("cinematic", "Cinematic film"), ("illustration", "Illustration or anime"),
    ("3d", "3D render"), ("sketch", "Pencil or sketch"), ("graphic", "Graphic design"),
])
LIGHTS = OrderedDict([
    ("daylight", "Daylight"), ("golden", "Golden hour"), ("night", "Night"), ("neon", "Neon or mixed"),
    ("practical", "Indoor practical"), ("fog-rain", "Fog or rain"), ("noon", "Harsh noon"),
])
LOCALES = OrderedDict([("IN", "Indian context"), ("global", "Global")])
SCRIPTS = OrderedDict([
    ("none", "No text"), ("latin", "Latin"), ("devanagari", "Devanagari"),
    ("telugu", "Telugu"), ("tamil", "Tamil"), ("mixed", "Mixed scripts"),
])
TEXT_SCRIPTS = ("latin", "devanagari", "telugu", "tamil")
SCRIPT_RANGES = {"devanagari": "ऀ-ॿ", "telugu": "ఀ-౿", "tamil": "஀-௿"}
DIFFS = OrderedDict([("hard", "Hard"), ("extreme", "Extreme")])

# Hardness levers: the image list plus four that only exist in motion.
LEVERS = OrderedDict([
    ("set", "Several clips that must agree with each other"),
    ("unseen", "Has to show what the input never shows (reverse side, back, behind the camera)"),
    ("count", "Exact numbers: people, objects, beats"),
    ("layout", "Exact positions: who is left or right, what is where"),
    ("exact-text", "Exact words that must appear letter for letter"),
    ("script", "Text in Devanagari, Telugu or Tamil"),
    ("small-text", "Small or dense text that must stay readable"),
    ("exclude", "Named things that must not appear"),
    ("off-centre", "Composition against the default centred framing"),
    ("camera", "A named lens, shot size, camera height or angle"),
    ("light-lock", "A light direction or time of day that must hold"),
    ("physics", "Reflections, shadows, refraction, occlusion or contact"),
    ("identity", "A face, character or product kept from a reference"),
    ("multi-ref", "Two to four inputs combined"),
    ("locale", "Indian cultural detail that a default model gets wrong"),
    ("keep-rest", "Everything outside the change stays the same"),
    ("multi-step", "Several changes in one instruction"),
    ("long-brief", "A long brief with ten or more constraints"),
    ("continuity", "Screen direction, props or costume carried across frames or clips"),
    ("time-code", "Time codes or durations that must add up"),
    ("camera-move", "A named camera move with a start and an end"),
    ("timing", "Actions that must happen at stated seconds"),
    ("end-frame", "The clip must start on one frame and end on another"),
    ("parallax", "Layers at different depths must move at different speeds"),
])

# Rubric criteria that only make sense with an input image or clip.
INPUT_CRITERIA = ("reference_fidelity", "preservation")

SET_MODES = OrderedDict([
    ("separate", "Separate calls, each from the same inputs"),
    ("chained", "Each clip starts from the previous clip's last frame"),
])
RUNNABLE = OrderedDict([("yes", "Yes"), ("partly", "Partly"), ("no", "No")])
I2V_ROLES = ("first_frame", "last_frame", "character", "environment", "product", "logo", "style",
             "still1", "still2", "still3", "still4")
EDIT_ROLES = ("source",) + I2V_ROLES

SAHIL = OrderedDict([
    ("composition", ("Composition and texture against a composition-led tool", "223–231", "VCAM")),
    ("product-bible", ("Production bible built from one product photo", "232–233", "VSPT")),
    ("storyboard", ("Storyboard built from a base image", "234–235", "VSTB")),
    ("beat-sheet", ("Time-coded beat sheet", "236–237", "VBTS")),
    ("typography", ("Typography posters against a type-led tool", "238–239", "VTXT")),
    ("angles", ("Many camera angles from one shot", "239–244", "VCOV")),
])

VENDOR_RE = re.compile(
    r"\b(gemini|nano[\s-]?banana|nb[\s-]?pro|imagen|veo|gpt|chatgpt|dall[\s·\-]?e|sora|midjourney|"
    r"recraft|flux|seedance|kling|runway[\s-]?(ml|gen)|stable[\s-]?diffusion|sdxl|firefly|ideogram|"
    r"leonardo[\s.]?ai|pika|hailuo|luma|openai|google|deepmind|bytedance|adobe|elevenlabs)\b",
    re.I)
TOOL_SYNTAX_RE = re.compile(
    r"(--[a-z]{1,12}\b|::\s*-?\d|\(\s*[\w\s-]+:\s*\d+(\.\d+)?\s*\)|\[\[|\]\]|\{\{|\}\})", re.I)

SLOT_ID_RE = re.compile(r"^(?P<fam>V[A-Z]{3})-(?P<ind>[A-Z]{3})-(?P<n>\d{2})$")
SOURCE_ID_RE = re.compile(r"^SRC-[A-Z0-9]+(-[A-Z0-9]+)*$")
ASPECT_RE = re.compile(r"^\d+(\.\d+)?:\d+$")


# ---------------------------------------------------------------- loading

def load_rubric(path: Path, lane: str = "video") -> dict:
    """The shared rubric: scoring rules, pass score and this lane's criteria by key."""
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    crit = OrderedDict()
    for c in data.get("criteria") or []:
        for k in ("key", "name", "lanes", "judges", "full", "half", "applies"):
            if not c.get(k):
                raise SystemExit(f"{path.name}: criterion {c.get('key', '?')} is missing {k}")
        if lane in c["lanes"]:
            crit[c["key"]] = c
    return {"scoring": data["scoring"], "pass_score": float(data["pass_score"]), "criteria": crit}


def load(src: Path):
    """Return (families by code, video sources by id)."""
    fams, parts = {}, []
    for f in sorted(src.glob("*.yaml")):
        if f.name.startswith("_") or f.name in ("sources.yaml", "customers.yaml"):
            continue
        data = yaml.safe_load(f.read_text(encoding="utf-8")) or {}
        data["_file"] = f.name
        if "family" in data:
            data["slots"] = list(data.get("slots") or [])
            data["scenarios"] = list(data.get("scenarios") or [])
            for s in data["scenarios"]:
                s["_file"] = f.name
            fams[data["code"]] = data
        else:
            parts.append(data)
    for p in parts:
        if p.get("code") not in fams:
            raise SystemExit(f"{p['_file']}: no family header for code {p.get('code')}")
        for s in p.get("scenarios") or []:
            s["_file"] = p["_file"]
            fams[p["code"]]["scenarios"].append(s)
    missing = [c for c in FAMILY_ORDER if c not in fams]
    if missing:
        raise SystemExit(f"missing family headers: {', '.join(missing)}")
    sources = OrderedDict()
    for s in (yaml.safe_load((src / "sources.yaml").read_text(encoding="utf-8")) or {}).get("sources") or []:
        if s.get("id") in sources:
            raise SystemExit(f"sources.yaml: duplicate source {s.get('id')}")
        sources[s["id"]] = s
    return OrderedDict((c, fams[c]) for c in FAMILY_ORDER), sources


def load_customers(path: Path, types_path: Path = None) -> tuple[dict, dict]:
    """(types by code, customer codes by scenario id) from customers.yaml. The
    types live in the image bank's file; the video bank reads them from there."""
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    types = (yaml.safe_load(types_path.read_text(encoding="utf-8")) if types_path else data).get("types") or {}
    for k, v in types.items():
        if not (v or {}).get("name") or not v.get("detail"):
            raise SystemExit(f"customers.yaml: type {k} needs a name and a detail")
    return types, data.get("scenarios") or {}


def validate_customers(r, written: dict, customers: tuple[dict, dict]) -> None:
    """Every scenario names 1–3 known customer types; no stray ids."""
    types, by_id = customers
    for sid in written:
        got = by_id.get(sid) or []
        if not 1 <= len(got) <= 3:
            r.err(f"{sid}: customers.yaml names {len(got)} target customers; a scenario has 1–3")
        for c in got:
            if c not in types:
                r.err(f"{sid}: unknown customer type '{c}' in customers.yaml")
        if len(set(got)) != len(got):
            r.err(f"{sid}: a target customer is named twice in customers.yaml")
    for sid in by_id:
        if sid not in written:
            r.err(f"customers.yaml: {sid} is not a scenario in the bank")


def customers_text(customers: tuple[dict, dict], sid: str) -> str:
    types, by_id = customers
    return ", ".join(types[c]["name"] for c in by_id[sid])


def customers_sheet(wb, scen: list, customers: tuple[dict, dict], table) -> None:
    types, by_id = customers
    ws = wb.create_sheet("Target customers")
    n = Counter(c for s in scen for c in by_id.get(s["id"], []))
    rows = [[v["name"], v["detail"], n[k]] for k, v in types.items() if n[k]]
    table(ws, ["Target customer", "Who it covers", "Scenarios"], rows, [32, 80, 10])


def load_image_bank(image_src: Path):
    """Image sources (read-only) and image slot ids with their family code."""
    sources = OrderedDict()
    data = yaml.safe_load((image_src / "sources.yaml").read_text(encoding="utf-8")) or {}
    for s in data.get("sources") or []:
        sources[s["id"]] = s
    slots = {}
    for f in sorted(image_src.glob("*.yaml")):
        if f.name == "sources.yaml":
            continue
        d = yaml.safe_load(f.read_text(encoding="utf-8")) or {}
        if "family" in d:
            for sl in d.get("slots") or []:
                slots[sl["id"]] = d["code"]
    return sources, slots


def words(text: str) -> int:
    return len(re.findall(r"\w+", text or ""))


def banned(text: str) -> list[str]:
    hits = [m.group(0) for m in VENDOR_RE.finditer(text or "")]
    hits += [m.group(0) for m in TOOL_SYNTAX_RE.finditer(text or "")]
    return hits


def has_script(text: str, script: str) -> bool:
    if script == "latin":
        return bool(re.search(r"[A-Za-z]", text))
    return bool(re.search(f"[{SCRIPT_RANGES[script]}]", text))


# ---------------------------------------------------------------- validation

class Report:
    def __init__(self):
        self.errors: list[str] = []
        self.rules: list[tuple[str, str, str, str]] = []

    def err(self, msg: str):
        self.errors.append(msg)

    def rule(self, scope: str, rule: str, value, status: str):
        self.rules.append((scope, rule, str(value), status))


def validate(fams, sources, img_sources, img_slots, rubric, customers) -> Report:
    r = Report()
    slots, seen = {}, set()

    # --- family rubrics
    for code, fam in fams.items():
        fr = fam.get("rubric") or {}
        w, gate = fr.get("weights") or {}, fr.get("gate") or []
        if not w or not gate or not fr.get("gate_rule"):
            r.err(f"{code}: the family rubric needs weights, a gate and a gate_rule")
        for k, v in w.items():
            if k not in rubric["criteria"]:
                r.err(f"{code}: rubric weight for unknown criterion '{k}'")
            if not (isinstance(v, int) and v > 0):
                r.err(f"{code}: rubric weight for {k} must be a whole number above 0")
        for k in gate:
            if k not in w:
                r.err(f"{code}: gate criterion '{k}' has no weight")
    for sid in sources:
        if sid in img_sources:
            r.err(f"{sid}: already an image source; reuse it by id instead of redefining it")
    all_sources = OrderedDict(list(img_sources.items()) + list(sources.items()))

    # --- slots
    for code, fam in fams.items():
        for t in fam.get("targets") or {}:
            if not re.match(r"^[a-z][a-z-]*$", t):
                r.err(f"{code}: target '{t}' must be lower-case words joined by hyphens")
        for sl in fam["slots"]:
            sid = sl.get("id", "?")
            m = SLOT_ID_RE.match(sid)
            if not m:
                r.err(f"{sid}: slot id must look like VCOV-FLM-01")
                continue
            if sid in seen:
                r.err(f"{sid}: duplicate slot id")
            seen.add(sid)
            if m["fam"] != code:
                r.err(f"{sid}: slot is in {fam['_file']} but its id says family {m['fam']}")
            if m["ind"] not in INDUSTRIES:
                r.err(f"{sid}: unknown industry {m['ind']}")
            elif m["ind"] not in PLAN[code]:
                r.err(f"{sid}: the plan has no {INDUSTRIES[m['ind']]} slots in {code}")
            for key, allowed in (("task", TASKS), ("subject", SUBJECTS), ("style", STYLES),
                                 ("light", LIGHTS), ("locale", LOCALES), ("diff", DIFFS)):
                if sl.get(key) not in allowed:
                    r.err(f"{sid}: {key} '{sl.get(key)}' is not one of {', '.join(allowed)}")
            sl.setdefault("script", "none")
            if sl["script"] not in SCRIPTS:
                r.err(f"{sid}: script '{sl['script']}' is not one of {', '.join(SCRIPTS)}")
            if fam.get("text_family") and sl["script"] == "none":
                r.err(f"{sid}: {code} is a text family, so every slot names its script")
            if sl.get("target") not in (fam.get("targets") or {}):
                r.err(f"{sid}: target '{sl.get('target')}' is not a {code} target")
            if not sl.get("idea"):
                r.err(f"{sid}: missing idea")
            if sl.get("from_image") and sl["from_image"] not in img_slots:
                r.err(f"{sid}: from_image {sl['from_image']} is not an image scenario")
            elif sl.get("from_image") and img_slots[sl["from_image"]] != BUILT_ON[code]:
                r.err(f"{sid}: from_image {sl['from_image']} is a {img_slots[sl['from_image']]} scenario; "
                      f"{code} is built on {BUILT_ON[code]}")
            srcs = sl.get("sources") or []
            sl["sources"] = srcs
            for s in srcs:
                if s not in all_sources:
                    r.err(f"{sid}: source {s} is in neither the video nor the image sources")
            kinds = Counter(all_sources[s].get("kind") for s in srcs if s in all_sources)
            task = sl.get("task")
            if task == "t2v" and srcs:
                r.err(f"{sid}: a text-to-video slot has no sources")
            if task == "i2v" and (not srcs or kinds["clip"]):
                r.err(f"{sid}: an image-to-video slot names one or more image sources and no clips")
            if task == "edit" and kinds["clip"] != 1:
                r.err(f"{sid}: a video edit names exactly one source clip")
            if len(srcs) > 4:
                r.err(f"{sid}: at most four inputs")
            sl["_fam"], sl["_ind"] = code, m["ind"]
            slots[sid] = sl

    # --- slot plan per family and industry
    for code, fam in fams.items():
        by_ind = defaultdict(list)
        for sl in fam["slots"]:
            by_ind[sl.get("_ind")].append(sl)
        for ind, (qty, t2v, i2v, edit) in PLAN[code].items():
            got = by_ind.get(ind, [])
            if not got:
                continue
            tasks = Counter(s["task"] for s in got)
            if len(got) != qty:
                r.err(f"{code} × {INDUSTRIES[ind]}: {len(got)} slots, the plan says {qty}")
            if (tasks["t2v"], tasks["i2v"], tasks["edit"]) != (t2v, i2v, edit):
                r.err(f"{code} × {INDUSTRIES[ind]}: task split {tasks['t2v']} t2v / {tasks['i2v']} i2v / "
                      f"{tasks['edit']} edit, the plan says {t2v} / {i2v} / {edit}")
            nums = sorted(int(s["id"][-2:]) for s in got)
            if nums != list(range(1, len(got) + 1)):
                r.err(f"{code} × {INDUSTRIES[ind]}: slot numbers must run 01–{len(got):02d} with no gaps")

    keys = defaultdict(list)
    for sl in slots.values():
        keys[(sl["_fam"], sl["subject"], sl["_ind"], sl["target"])].append(sl["id"])
    for k, ids in keys.items():
        if len(ids) > 1:
            r.err(f"{', '.join(ids)}: same family, subject, industry and failure target ({k[1]}, {k[3]})")

    for code, fam in fams.items():
        planned = sum(v[0] for v in PLAN[code].values())
        subj = Counter(sl["subject"] for sl in fam["slots"])
        cap = int(planned * MAX_SUBJECT_SHARE)
        for s, n in subj.items():
            if n > cap:
                r.err(f"{code}: {n} '{s}' slots; no subject may pass {cap} ({MAX_SUBJECT_SHARE:.0%} of {planned})")
        diffs = Counter(sl["diff"] for sl in fam["slots"])
        if diffs["extreme"] > EXTREME[code]:
            r.err(f"{code}: {diffs['extreme']} extreme slots, the plan allows {EXTREME[code]}")
        if diffs["hard"] > planned - EXTREME[code]:
            r.err(f"{code}: {diffs['hard']} hard slots, the plan allows {planned - EXTREME[code]}")

    # --- sources (video-only) and uses across both lanes
    uses = defaultdict(list)
    for sl in slots.values():
        for s in sl["sources"]:
            uses[s].append(sl["id"])
    for sid, s in sources.items():
        if not SOURCE_ID_RE.match(sid):
            r.err(f"{sid}: source ids look like SRC-KIDS-TREE")
        if s.get("kind") not in ("image", "clip"):
            r.err(f"{sid}: kind must be image or clip")
        if s.get("role") not in ("scene", "reference"):
            r.err(f"{sid}: role must be scene or reference")
        for k in ("source_prompt", "aspect", "resolution"):
            if not s.get(k):
                r.err(f"{sid}: missing {k}")
        if not ASPECT_RE.match(str(s.get("aspect", ""))):
            r.err(f"{sid}: aspect looks like 16:9")
        if s.get("kind") == "clip":
            d = s.get("duration")
            if not isinstance(d, (int, float)) or not DURATION_RANGE[0] <= d <= DURATION_RANGE[1]:
                r.err(f"{sid}: a clip states its duration in seconds ({DURATION_RANGE[0]}–{DURATION_RANGE[1]})")
            if not s.get("camera_move") or not s.get("action_timing"):
                r.err(f"{sid}: a clip states its camera move and its action timing")
        if words(s.get("source_prompt", "")) < MIN_SOURCE_PROMPT_WORDS:
            r.err(f"{sid}: source prompt under {MIN_SOURCE_PROMPT_WORDS} words; describe layout, light and lens")
        if len(s.get("must_contain") or []) < 2:
            r.err(f"{sid}: must_contain needs at least two acceptance items")
        if not s.get("avoid"):
            r.err(f"{sid}: avoid needs at least one item")
        text = " ".join([s.get("source_prompt", ""), str(s.get("camera_move", "")), str(s.get("action_timing", ""))]
                        + list(s.get("avoid") or []) + list(s.get("must_contain") or []))
        hits = banned(text)
        if hits:
            r.err(f"{sid}: source prompt names a model or vendor, or uses tool syntax: {', '.join(sorted(set(hits)))}")
        declared = sorted(s.get("used_by") or [])
        if declared != sorted(uses.get(sid, [])):
            r.err(f"{sid}: used_by {declared} does not match the slots that name it {sorted(uses.get(sid, []))}")
    for sid, users in uses.items():
        img_uses = len((img_sources.get(sid) or {}).get("used_by") or [])
        if img_uses + len(users) > MAX_KIT_USES:
            r.err(f"{sid}: {img_uses} image + {len(users)} video uses; the cap is {MAX_KIT_USES}")

    # --- scenarios
    written = {}
    for code, fam in fams.items():
        for s in fam["scenarios"]:
            sid = s.get("id", "?")
            where = f"{sid} ({s.get('_file')})"
            if sid not in slots:
                r.err(f"{where}: no slot with this id")
                continue
            sl = slots[sid]
            if sl["_fam"] != code:
                r.err(f"{where}: filed under {code} but the slot is {sl['_fam']}")
            if sid in written:
                r.err(f"{where}: written twice")
            written[sid] = s
            s["_slot"] = sl
            validate_scenario(r, where, s, sl, fam, all_sources, rubric)

    coverage(r, fams, slots, written, sources, uses, img_sources)
    validate_customers(r, written, customers)
    return r


def validate_scenario(r, where, s, sl, fam, all_sources, rubric) -> None:
    for k in ("title", "persona", "evidence", "levers", "prompt", "clip", "deliverable", "checks",
              "competitor", "runnable_today"):
        if not s.get(k):
            r.err(f"{where}: missing {k}")
    levers = s.get("levers") or []
    for lv in levers:
        if lv not in LEVERS:
            r.err(f"{where}: unknown lever '{lv}'")
    if len(set(levers)) < MIN_LEVERS:
        r.err(f"{where}: {len(set(levers))} levers; each scenario uses at least {MIN_LEVERS}")
    if words(s.get("prompt", "")) < MIN_PROMPT_WORDS:
        r.err(f"{where}: prompt under {MIN_PROMPT_WORDS} words")

    clip = s.get("clip") or {}
    d = clip.get("duration")
    if not isinstance(d, (int, float)) or not DURATION_RANGE[0] <= d <= DURATION_RANGE[1]:
        r.err(f"{where}: clip.duration is the length in seconds ({DURATION_RANGE[0]}–{DURATION_RANGE[1]})")
    if not ASPECT_RE.match(str(clip.get("aspect", ""))):
        r.err(f"{where}: clip.aspect looks like 16:9")
    if not clip.get("camera"):
        r.err(f"{where}: clip.camera names the camera move (or 'locked off')")

    inputs = s.get("inputs") or []
    task = sl["task"]
    got = sorted(i.get("source") for i in inputs)
    if got != sorted(sl["sources"]):
        r.err(f"{where}: inputs {got} do not match the slot's sources {sorted(sl['sources'])}")
    roles = Counter(i.get("role") for i in inputs)
    if task == "i2v":
        if any(x not in I2V_ROLES for x in roles):
            r.err(f"{where}: image-to-video roles are {', '.join(I2V_ROLES)}")
        if roles["first_frame"] > 1 or roles["last_frame"] > 1:
            r.err(f"{where}: at most one first_frame and one last_frame")
        if roles["last_frame"] and not roles["first_frame"]:
            r.err(f"{where}: a last_frame needs a first_frame")
    elif task == "edit":
        if any(x not in EDIT_ROLES for x in roles):
            r.err(f"{where}: video-edit roles are source plus {', '.join(I2V_ROLES)}")
        if roles["source"] != 1:
            r.err(f"{where}: a video edit has exactly one input with role source")
        for i in inputs:
            if i.get("role") == "source" and (all_sources.get(i.get("source")) or {}).get("kind") != "clip":
                r.err(f"{where}: the source of a video edit is a clip")
    elif inputs:
        r.err(f"{where}: a text-to-video scenario has no inputs")

    checks = s.get("checks") or []
    lo, hi = CHECKS_RANGE
    if not lo <= len(checks) <= hi:
        r.err(f"{where}: {len(checks)} checks; a scenario has {lo}–{hi}")
    crits = Counter()
    for i, c in enumerate(checks, 1):
        if not (isinstance(c, list) and len(c) == 2):
            r.err(f"{where}: check {i} must be [criterion, question]")
            continue
        if c[0] not in rubric["criteria"]:
            r.err(f"{where}: check {i} has unknown rubric criterion '{c[0]}'")
        if not str(c[1]).rstrip().endswith("?"):
            r.err(f"{where}: check {i} must be a yes/no question ending in '?'")
        crits[c[0]] += 1
    if not (crits["motion_coherence"] or crits["temporal_consistency"]):
        r.err(f"{where}: a video scenario checks motion coherence or temporal consistency at least once")

    sm = s.get("set_members")
    rng = fam.get("set")
    all_prompts = [s.get("prompt", "")]
    if rng and (sm or rng.get("required")):
        if not sm:
            r.err(f"{where}: {sl['_fam']} scenarios are sets; add set_members")
        else:
            if sm.get("mode") not in SET_MODES:
                r.err(f"{where}: set mode must be one of {', '.join(SET_MODES)}")
            if len(sm.get("must_stay_same") or []) < 3:
                r.err(f"{where}: must_stay_same needs at least three items")
            members = sm.get("members") or []
            if not rng["min"] <= len(members) <= rng["max"]:
                r.err(f"{where}: {len(members)} clips in the set; {sl['_fam']} sets have {rng['min']}–{rng['max']}")
            mlo, mhi = MEMBER_CHECKS_RANGE
            for j, mem in enumerate(members, 1):
                if not mem.get("name") or not mem.get("prompt"):
                    r.err(f"{where}: clip {j} needs a name and a prompt")
                all_prompts.append(mem.get("prompt", ""))
                mc = mem.get("checks") or []
                if not mlo <= len(mc) <= mhi:
                    r.err(f"{where}: clip {j} has {len(mc)} checks; clips have {mlo}–{mhi}")
                for k, c in enumerate(mc, 1):
                    if not (isinstance(c, list) and len(c) == 2 and c[0] in rubric["criteria"]
                            and str(c[1]).rstrip().endswith("?")):
                        r.err(f"{where}: clip {j} check {k} must be [criterion, yes/no question?]")
                    else:
                        crits[c[0]] += 1
            if not crits["set_consistency"]:
                r.err(f"{where}: a set needs at least one set_consistency check")
            if "set" not in levers:
                r.err(f"{where}: a set scenario uses the 'set' lever")
    elif sm:
        r.err(f"{where}: {sl['_fam']} scenarios are single clips; remove set_members")

    hits = banned(" ".join(all_prompts + [str(clip.get("camera", ""))]))
    if hits:
        r.err(f"{where}: prompt names a model or vendor, or uses tool syntax: {', '.join(sorted(set(hits)))}")

    exact = s.get("exact_text") or []
    joined = " ".join(" ".join(all_prompts).split())
    for t in exact:
        if " ".join(str(t).split()) not in joined:
            r.err(f"{where}: exact text '{t}' is not given word for word in the prompt")
    if exact and not crits["text_accuracy"]:
        r.err(f"{where}: exact text needs at least one text_accuracy check")

    # rubric: every criterion checked is weighted by the family, and the gate applies
    fr = fam.get("rubric") or {}
    if not crits["prompt_adherence"]:
        r.err(f"{where}: every scenario has at least one prompt_adherence check")
    for k in crits:
        if k in rubric["criteria"] and k not in (fr.get("weights") or {}):
            r.err(f"{where}: checks {k}, which the {sl['_fam']} rubric does not weight")
        if k in INPUT_CRITERIA and not inputs:
            r.err(f"{where}: {k} checks need an input image or clip")
    if fr.get("gate") and not any(crits[k] for k in fr["gate"]):
        r.err(f"{where}: checks none of the gate criteria ({', '.join(fr['gate'])})")
    s["_crits"] = crits
    script = sl["script"]
    if script in SCRIPT_RANGES and not any(has_script(t, script) for t in exact):
        r.err(f"{where}: the slot's script is {script} but no exact text uses it")
    if script == "mixed" and sum(any(has_script(t, sc) for t in exact) for sc in TEXT_SCRIPTS) < 2:
        r.err(f"{where}: script is mixed but the exact text uses fewer than two scripts")
    if script == "none" and exact:
        r.err(f"{where}: the slot has no text but the scenario lists exact text; set the slot's script")

    rt = s.get("runnable_today") or {}
    if isinstance(rt.get("status"), bool):
        r.err(f"{where}: write runnable_today status in quotes (\"yes\"); YAML reads a bare yes as true")
    elif rt.get("status") not in RUNNABLE or not rt.get("note"):
        r.err(f"{where}: runnable_today needs status (yes, partly, no) and a note")
    if s.get("sahil") and s["sahil"] not in SAHIL:
        r.err(f"{where}: sahil must be one of {', '.join(SAHIL)}")
    if s.get("sahil") and not s.get("seed"):
        r.err(f"{where}: a recreation of Sahil's example is a seed")


def coverage(r, fams, slots, written, sources, uses, img_sources) -> None:
    all_slots = list(slots.values())
    bank_full = len(all_slots) == BANK_TOTAL
    for code, fam in fams.items():
        planned = sum(v[0] for v in PLAN[code].values())
        fs = fam["slots"]
        full = len(fs) == planned
        scope = f"{code} {fam['family']}"
        st = lambda ok: ("met" if ok else "broken") if full else "pending"

        def chk(rule, value, ok):
            r.rule(scope, rule, value, st(ok))
            if full and not ok:
                r.err(f"{code}: {rule} — {value}")

        chk("Slots declared", f"{len(fs)} of {planned}", len(fs) == planned)
        inds = {s["_ind"] for s in fs}
        chk(f"At least {MIN_PER_FAMILY['industries']} industries", len(inds), len(inds) >= MIN_PER_FAMILY["industries"])
        subj = Counter(s["subject"] for s in fs)
        chk(f"At least {MIN_PER_FAMILY['subjects']} subjects", len(subj), len(subj) >= MIN_PER_FAMILY["subjects"])
        sty = {s["style"] for s in fs}
        chk(f"At least {MIN_PER_FAMILY['styles']} styles", len(sty), len(sty) >= MIN_PER_FAMILY["styles"])
        lig = {s["light"] for s in fs}
        chk(f"At least {MIN_PER_FAMILY['lights']} light conditions", len(lig), len(lig) >= MIN_PER_FAMILY["lights"])
        top = subj.most_common(1)[0] if subj else ("—", 0)
        share = top[1] / max(len(fs), 1)
        chk(f"No subject above {MAX_SUBJECT_SHARE:.0%}", f"{top[0]} {share:.0%}", share <= MAX_SUBJECT_SHARE)
        ext = sum(1 for s in fs if s["diff"] == "extreme")
        chk("Extreme share about 40%", f"{ext} of {len(fs)} (plan {EXTREME[code]})", ext == EXTREME[code])
        t = Counter(s["task"] for s in fs)
        pt = tuple(sum(v[i] for v in PLAN[code].values()) for i in (1, 2, 3))
        chk("Task split matches the plan",
            f"{t['t2v']} t2v / {t['i2v']} i2v / {t['edit']} edit (plan {pt[0]} / {pt[1]} / {pt[2]})",
            (t["t2v"], t["i2v"], t["edit"]) == pt)
        if fam.get("text_family"):
            scr = {s["script"] for s in fs}
            chk("All four scripts", ", ".join(sorted(scr)) or "—", all(x in scr for x in TEXT_SCRIPTS))
        n_written = sum(1 for s in fs if s["id"] in written)
        r.rule(scope, "Scenarios written", f"{n_written} of {planned}", "met" if n_written == planned else "pending")
        seeds = [s for s in fam["scenarios"] if s.get("seed")]
        r.rule(scope, "Seeds (3 per family)", len(seeds), "met" if len(seeds) == 3 else ("broken" if len(seeds) > 3 else "pending"))
        if len(seeds) > 3:
            r.err(f"{code}: {len(seeds)} seeds; a family has 3")

    scope = "Whole bank"
    n = len(all_slots)
    india = sum(1 for s in all_slots if s["locale"] == "IN")
    ok = india / max(n, 1) >= MIN_INDIA
    r.rule(scope, f"At least {MIN_INDIA:.0%} Indian context", f"{india} of {n} ({india / max(n, 1):.0%})",
           ("met" if ok else "broken") if bank_full else ("on track" if ok or not n else "behind"))
    if bank_full and not ok:
        r.err(f"bank: Indian context {india / n:.0%}, needs {MIN_INDIA:.0%}")
    tasks = Counter(s["task"] for s in all_slots)
    for task, share in MIN_SHARE.items():
        got = tasks[task] / max(n, 1)
        ok = got >= share
        r.rule(scope, f"{TASKS[task]} at least {share:.0%}", f"{tasks[task]} of {n} ({got:.0%})",
               ("met" if ok else "broken") if bank_full else ("on track" if ok or not n else "behind"))
        if bank_full and not ok:
            r.err(f"bank: {TASKS[task]} share {got:.0%}, needs {share:.0%}")
    all_written = bank_full and len(written) == BANK_TOTAL
    lever_fams = defaultdict(set)
    for s in written.values():
        for lv in s.get("levers") or []:
            lever_fams[lv].add(s["_slot"]["_fam"])
    thin = [lv for lv in LEVERS if lv not in LEVER_FAMILY_EXEMPT and len(lever_fams[lv]) < MIN_LEVER_FAMILIES]
    r.rule(scope, f"Every lever in at least {MIN_LEVER_FAMILIES} families (except: "
           + "; ".join(f"{k}, {v}" for k, v in LEVER_FAMILY_EXEMPT.items()) + ")",
           "all" if not thin else f"{len(thin)} below: {', '.join(thin)}",
           ("met" if not thin else "broken") if all_written else "pending")
    if all_written and thin:
        r.err(f"bank: levers in fewer than {MIN_LEVER_FAMILIES} families: {', '.join(thin)}")
    unused = [sid for sid in sources if not uses.get(sid)]
    r.rule(scope, "Every video source used by a slot", "all" if not unused else ", ".join(unused),
           "met" if not unused else "broken")
    for sid in unused:
        r.err(f"{sid}: no slot uses this source")
    reused = sum(1 for sid in uses if sid in img_sources)
    r.rule(scope, "Image sources reused in video", reused, "met" if reused else "pending")
    sahil_done = {s.get("sahil") for s in written.values() if s.get("sahil")}
    r.rule(scope, "Sahil's six examples as seeds", f"{len(sahil_done)} of 6",
           "met" if len(sahil_done) == 6 else "pending")


# ---------------------------------------------------------------- derived values

def outputs_of(s: dict) -> int:
    sm = s.get("set_members")
    return len(sm.get("members") or []) if sm else 1


def seconds_of(s: dict) -> float:
    """clip.duration is one clip's length; for a chained set it is the whole passage."""
    d = (s.get("clip") or {}).get("duration", 0)
    sm = s.get("set_members") or {}
    return d if sm.get("mode") == "chained" else outputs_of(s) * d


def all_checks(s: dict) -> list:
    return list(s["checks"]) + [c for m in (s.get("set_members") or {}).get("members") or []
                                for c in m.get("checks") or []]


def members_text(s: dict) -> str:
    sm = s.get("set_members") or {}
    return "\n".join(f"{j}. {m['name']}: {m['prompt'].strip()}"
                     for j, m in enumerate(sm.get("members") or [], 1)) or "—"


def rubric_of(s: dict, fam: dict) -> list:
    """[(criterion, weight share, checks)] for the criteria this scenario checks."""
    w = fam["rubric"]["weights"]
    crits = s["_crits"]
    total = sum(w[k] for k in crits)
    return [(k, w[k] / total, crits[k]) for k in w if k in crits]


def rubric_text(s: dict, fam: dict, rubric: dict) -> str:
    gate = fam["rubric"]["gate"]
    return "\n".join(f"{rubric['criteria'][k]['name']}: {share:.0%} · {n} check{'s' if n > 1 else ''}"
                     + (" · gate" if k in gate else "") for k, share, n in rubric_of(s, fam))


def pass_text(s: dict, fam: dict, rubric: dict) -> str:
    gate = [rubric["criteria"][k]["name"] for k in fam["rubric"]["gate"] if s["_crits"][k]]
    n = sum(s["_crits"][k] for k in fam["rubric"]["gate"])
    first = (f"The {gate[0]} check is a yes" if n == 1
             else f"All {n} {' and '.join(gate)} checks are yes")
    return f"{first}, and the score is {rubric['pass_score']:.1f} or more out of 10."


def target_label(fam: dict, t: str) -> str:
    return f"{t} — {(fam.get('targets') or {}).get(t, '')}"


def inputs_text(s: dict) -> str:
    return "\n".join(f"{i['source']} ({i['role']})" + (f": {i['note']}" if i.get("note") else "")
                     for i in s.get("inputs") or []) or "—"


# ---------------------------------------------------------------- workbook helpers

HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
SECTION_FONT = Font(bold=True, size=11, color="1F3A5F")
WRAP = Alignment(wrap_text=True, vertical="top")
STATUS_FILL = {
    "met": PatternFill("solid", fgColor="D9EAD3"), "on track": PatternFill("solid", fgColor="D9EAD3"),
    "pending": PatternFill("solid", fgColor="FFF2CC"), "behind": PatternFill("solid", fgColor="FCE5CD"),
    "broken": PatternFill("solid", fgColor="F4CCCC"),
}


def table(ws, headers, rows, widths, freeze="B2"):
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, WRAP
    for row in rows:
        ws.append(row)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions


# ---------------------------------------------------------------- sheets

def build(fams, sources, img_sources, rubric, rep: Report, customers) -> Workbook:
    slots = [sl for f in fams.values() for sl in f["slots"]]
    scen = [s for f in fams.values() for s in sorted(f["scenarios"], key=lambda x: x["id"])]
    all_sources = OrderedDict(list(img_sources.items()) + list(sources.items()))
    wb = Workbook()

    ws = wb.active
    ws.title = "How to read"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 120
    ws.append(["Video scenario bank v3 — 240 scenarios, written to the 98th–99th percentile bar"])
    ws["A1"].font = TITLE_FONT
    outs = sum(outputs_of(s) for s in scen)
    secs = sum(seconds_of(s) for s in scen)
    started = [INDUSTRIES[i] for i in INDUSTRIES if any(sl["_ind"] == i for sl in slots)]
    paras = [
        ("What this is", "Hard text-to-video, image-to-video and video-edit scenarios across 9 families and 6 "
         "industries, each built on an image family so the image-to-video chain is part of the test. One "
         "scenario is one brief and one answer key: a four-clip coverage set is one scenario with four clips."),
        ("Progress", f"{len(slots)} of {BANK_TOTAL} slots declared; {len(scen)} scenarios written "
         f"({sum(1 for s in scen if s.get('seed'))} seeds), {outs} clips and {secs:g} seconds per model so far. "
         f"Industries started: {', '.join(started) or 'none'}. The bank is written one industry at a time."),
        ("Inputs", "Image-to-video scenarios reuse the image bank's sources as first frames and references; "
         "video edits use source clips with their own model-neutral prompts (Sources tab). A base scene's uses "
         "across both banks are capped at 5."),
        ("Sets", "A set is one scenario with several clips. Each clip is a separate request: the shared Prompt "
         "plus one line from 'What each clip is'. In a chained set each clip starts from the last frame of the "
         "clip before. The set is judged as a whole, so the clips must agree."),
        ("Target customers", "The kinds of business each scenario is for: where its persona works, then the "
         "brand category it serves. One to three per scenario, from a fixed list (Target customers tab)."),
        ("Review tab", "Persona, failure target and levers are for reviewing the bank, not for running it, "
         "so they sit on the Review tab."),
        ("Rubric", f"Every scenario is scored against one rubric of {len(rubric['criteria'])} criteria (Rubric tab), "
         "shared with the image bank plus motion coherence and temporal consistency. "
         f"Each scenario has {CHECKS_RANGE[0]}–{CHECKS_RANGE[1]} yes/no checks on the whole output, at least one "
         f"on motion coherence or temporal consistency, and each clip in a set {MEMBER_CHECKS_RANGE[0]}–{MEMBER_CHECKS_RANGE[1]} more; every check "
         "belongs to one criterion (Checks tab). " + rubric["scoring"]["criterion"] + " "
         + rubric["scoring"]["scenario"] + " No lip-sync and no audio are judged in this bank."),
        ("Pass", rubric["scoring"]["gate"] + " " + rubric["scoring"]["pass"]
         + f" The pass score is {rubric['pass_score']:.1f}. Each family's weights and gate are on the "
         "Rubric by family tab; each scenario's own weights and pass rule are on the Scenario Bank tab."),
        ("Sources", "Never use a model in the comparison to make its own inputs. Which tool made each source "
         "is recorded in a provenance file when it is created, not in this bank."),
        ("Coverage", "Rules that need a full family or a full bank show as pending until then; the build "
         "stops on any rule that is already broken."),
    ]
    for head, body in paras:
        ws.append([head, body])
        ws.cell(ws.max_row, 1).font = SECTION_FONT
        ws.cell(ws.max_row, 2).alignment = WRAP
    ws.column_dimensions["C"].width = 60
    ws.append([])
    ws.append(["Hardness lever", "What it means", "Example scenarios"])
    for c in ws[ws.max_row]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    shown = Counter()
    for k, v in LEVERS.items():
        ex = [s for s in scen if k in s["levers"]]
        ex = sorted(ex, key=lambda s: (not s.get("seed"), shown[s["id"]], s["id"]))[:2]
        shown.update(s["id"] for s in ex)
        ws.append([k, v, "\n".join(f"{s['id']}: {s['title']}" for s in ex)])
        ws.cell(ws.max_row, 3).alignment = WRAP
    ws.append([])
    ws.append(["Set mode", "What it means"])
    for c in ws[ws.max_row]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    used = {(s.get("set_members") or {}).get("mode") for s in scen}
    for k, v in SET_MODES.items():
        if k in used:
            ws.append([k, v])

    ws = wb.create_sheet("Scenario Bank")
    rows = []
    for s in scen:
        sl, fam = s["_slot"], fams[s["_slot"]["_fam"]]
        sm = s.get("set_members") or {}
        clip = s.get("clip") or {}
        rows.append([
            s["id"], fam["family"], INDUSTRIES[sl["_ind"]], customers_text(customers, s["id"]),
            TASKS[sl["task"]], DIFFS[sl["diff"]], s["title"], inputs_text(s), s["prompt"].strip(),
            clip.get("duration"), clip.get("aspect"), clip.get("camera"),
            SET_MODES.get(sm.get("mode"), "—") if sm else "—", outputs_of(s), members_text(s), s["deliverable"],
            "\n".join(s.get("exact_text") or []) or "—", sum(s["_crits"].values()),
            rubric_text(s, fam, rubric), pass_text(s, fam, rubric),
            sl.get("from_image") or "—",
            SUBJECTS[sl["subject"]], STYLES[sl["style"]], LIGHTS[sl["light"]], LOCALES[sl["locale"]],
            SCRIPTS[sl["script"]],
        ])
    table(ws, ["Scenario ID", "Family", "Industry", "Target customers", "Task", "Difficulty", "Title",
               "Inputs (source id, role)",
               "Prompt", "Seconds", "Aspect", "Camera", "Set mode", "Clips", "What each clip is", "Deliverable",
               "Exact text", "Checks", "Rubric (weight · checks)", "Pass rule", "Built on image scenario",
               "Subject", "Style", "Light", "Locale", "Script"],
          rows, [14, 20, 20, 30, 12, 10, 30, 30, 90, 8, 8, 24, 18, 6, 70, 34, 30, 8, 40, 50, 14,
                 14, 14, 14, 12, 12])

    ws = wb.create_sheet("Review")
    rows = [[s["id"], s["title"], s["persona"], target_label(fams[s["_slot"]["_fam"]], s["_slot"]["target"]),
             ", ".join(s["levers"])] for s in scen]
    table(ws, ["Scenario ID", "Title", "Persona", "Failure target", "Levers"], rows, [14, 40, 50, 50, 40])
    customers_sheet(wb, scen, customers, table)

    ws = wb.create_sheet("Rubric")
    used = Counter(c[0] for s in scen for c in all_checks(s))
    rows = [[c["name"], k, c["judges"], c["full"], c["half"], c["applies"], used[k]]
            for k, c in rubric["criteria"].items()]
    table(ws, ["Criterion", "Key", "What it judges", "Full marks (10) look like", "Half marks (5) look like",
               "Scored in", "Checks in the bank"], rows, [22, 20, 60, 50, 60, 34, 10])

    ws = wb.create_sheet("Rubric by family")
    rows = []
    for code, f in fams.items():
        fr = f["rubric"]
        fc = Counter(c[0] for s in f["scenarios"] for c in all_checks(s))
        rows.append([f["family"], "\n".join(f"{rubric['criteria'][k]['name']}: {v}" for k, v in fr["weights"].items()),
                     ", ".join(rubric["criteria"][k]["name"] for k in fr["gate"]), fr["gate_rule"],
                     "\n".join(f"{rubric['criteria'][k]['name']}: {fc[k]}" for k in fr["weights"] if fc[k])])
    table(ws, ["Family", "Weights (points; shares are taken over the criteria a scenario checks)",
               "Gate", "Why it is the gate", "Checks per criterion"], rows, [26, 40, 30, 50, 34])

    ws = wb.create_sheet("Checks")
    rows = []
    for s in scen:
        gate = fams[s["_slot"]["_fam"]]["rubric"]["gate"]
        for i, (k, q) in enumerate(s["checks"], 1):
            rows.append([s["id"], "Whole output", i, rubric["criteria"][k]["name"], "Yes" if k in gate else "", q])
        for mem in (s.get("set_members") or {}).get("members") or []:
            for i, (k, q) in enumerate(mem.get("checks") or [], 1):
                rows.append([s["id"], mem["name"], i, rubric["criteria"][k]["name"], "Yes" if k in gate else "", q])
    table(ws, ["Scenario ID", "Applies to", "#", "Rubric criterion", "Gate", "Yes/no check"],
          rows, [14, 30, 5, 22, 7, 110])

    ws = wb.create_sheet("Set members")
    rows = []
    for s in scen:
        sm = s.get("set_members")
        if not sm:
            continue
        rows.append([s["id"], "Must stay the same", "", "\n".join(f"• {x}" for x in sm["must_stay_same"]), ""])
        for j, mem in enumerate(sm["members"], 1):
            rows.append([s["id"], f"{j}. {mem['name']}", SET_MODES[sm["mode"]], mem["prompt"].strip(),
                         "\n".join(f"{rubric['criteria'][d]['name']}: {q}" for d, q in mem["checks"])])
    table(ws, ["Scenario ID", "Clip", "Set mode", "Clip brief (sent with the shared prompt)", "Clip checks"],
          rows, [14, 30, 26, 80, 70])

    ws = wb.create_sheet("Slot grid")
    written = {s["id"] for s in scen}
    rows = []
    for f in fams.values():
        for sl in f["slots"]:
            rows.append([sl["id"], f["family"], INDUSTRIES[sl["_ind"]], TASKS[sl["task"]], DIFFS[sl["diff"]],
                         SUBJECTS[sl["subject"]], STYLES[sl["style"]], LIGHTS[sl["light"]],
                         LOCALES[sl["locale"]], SCRIPTS[sl["script"]], target_label(f, sl["target"]),
                         ", ".join(sl["sources"]) or "—", sl.get("from_image") or "—", sl["idea"],
                         "Written" if sl["id"] in written else "Empty"])
    table(ws, ["Slot", "Family", "Industry", "Task", "Difficulty", "Subject", "Style", "Light", "Locale",
               "Script", "Failure target", "Sources", "Built on image scenario", "Idea", "Status"],
          rows, [14, 20, 20, 12, 10, 16, 16, 15, 13, 12, 40, 30, 14, 70, 9])

    ws = wb.create_sheet("Scene kits")
    users = defaultdict(list)
    for sl in slots:
        for sid in sl["sources"]:
            users[sid].append(sl["id"])
    rows = []
    for sid, src in all_sources.items():
        if src.get("role") != "scene" or (sid in img_sources and not users.get(sid)):
            continue
        img_users = list((img_sources.get(sid) or {}).get("used_by") or [])
        vid_users = users.get(sid, [])
        rows.append([sid, src.get("kind"), "image bank" if sid in img_sources else "video bank",
                     len(img_users), len(vid_users), MAX_KIT_USES - len(img_users) - len(vid_users),
                     ", ".join(img_users) or "—", ", ".join(vid_users) or "—"])
    table(ws, ["Scene kit", "Kind", "Defined in", "Image uses", "Video uses", "Uses left (cap 5)",
               "Image scenarios", "Video scenarios"], rows, [22, 7, 12, 8, 8, 10, 50, 50])

    ws = wb.create_sheet("Sources")
    rows = []
    for sid, src in sources.items():
        rows.append([sid, src["kind"], src["role"], ", ".join(src.get("used_by") or []), src["aspect"],
                     src["resolution"], src.get("duration") or "—", src.get("camera_move") or "—",
                     src.get("action_timing") or "—", src["source_prompt"].strip(),
                     "\n".join(f"• {x}" for x in src.get("avoid") or []),
                     "\n".join(f"☐ {x}" for x in src.get("must_contain") or []), src.get("alt_real") or "—"])
    table(ws, ["Source", "Kind", "Role", "Used by", "Aspect", "Resolution", "Seconds", "Camera move",
               "Action timing", "Source prompt (model-neutral)", "Must not appear", "Acceptance checklist",
               "Real footage instead"],
          rows, [22, 7, 10, 34, 8, 18, 8, 24, 40, 100, 34, 44, 34])

    ws = wb.create_sheet("Image sources reused")
    rows = [[sid, img_sources[sid].get("kind"), img_sources[sid].get("role"),
             len(img_sources[sid].get("used_by") or []), ", ".join(us)]
            for sid, us in users.items() if sid in img_sources]
    table(ws, ["Image source", "Kind", "Role", "Image uses", "Video slots"], rows, [24, 7, 10, 10, 60])

    ws = wb.create_sheet("Task-type split")
    rows = []
    for code, f in fams.items():
        for ind, (qty, t2v, i2v, edit) in PLAN[code].items():
            got = [sl for sl in f["slots"] if sl["_ind"] == ind]
            t = Counter(sl["task"] for sl in got)
            rows.append([f["family"], INDUSTRIES[ind], qty, t2v, i2v, edit, len(got), t["t2v"], t["i2v"],
                         t["edit"], sum(1 for sl in got if sl["id"] in written)])
    tt = Counter(sl["task"] for sl in slots)
    rows.append(["Total", "", BANK_TOTAL] + [sum(v[i] for p in PLAN.values() for v in p.values()) for i in (1, 2, 3)]
                + [len(slots), tt["t2v"], tt["i2v"], tt["edit"], len(scen)])
    table(ws, ["Family", "Industry", "Planned", "Planned t2v", "Planned i2v", "Planned edit", "Slots declared",
               "Declared t2v", "Declared i2v", "Declared edit", "Written"],
          rows, [26, 34, 9, 10, 10, 10, 10, 10, 10, 10, 9])

    ws = wb.create_sheet("Families")
    rows = []
    for code, f in fams.items():
        planned = sum(v[0] for v in PLAN[code].values())
        sah = [f"{v[0]} ({v[1]})" for k, v in SAHIL.items() if v[2] == code]
        rows.append([f["number"], code, f["family"], BUILT_ON[code], f["about"], planned, f["outputs"],
                     *[sum(v[i] for v in PLAN[code].values()) for i in (1, 2, 3)],
                     "\n".join(f"{k} — {v}" for k, v in (f.get("targets") or {}).items()),
                     "\n".join(sah) or "—",
                     ", ".join(s["id"] for s in f["scenarios"] if s.get("seed")) or "—"])
    table(ws, ["#", "Code", "Family", "Built on", "What it tests", "Scenarios", "Clips each", "t2v", "i2v",
               "edit", "Failure targets", "Sahil's examples (screenshots)", "Seeds written"],
          rows, [5, 7, 24, 8, 50, 9, 14, 6, 6, 6, 70, 40, 30])

    ws = wb.create_sheet("Coverage matrix")
    table(ws, ["Scope", "Rule", "Value", "Status"], [list(x) for x in rep.rules], [34, 40, 50, 11])
    for row in ws.iter_rows(min_row=2):
        fill = STATUS_FILL.get(row[3].value)
        if fill:
            row[3].fill = fill
    return wb


# ---------------------------------------------------------------- main

def progress_table(fams) -> str:
    lines = ["family   " + " ".join(f"{i:>9}" for i in INDUSTRIES) + "     total"]
    for code, f in fams.items():
        cells = []
        for ind in INDUSTRIES:
            if ind not in PLAN[code]:
                cells.append(f"{'·':>9}")
                continue
            n_slots = sum(1 for sl in f["slots"] if sl["_ind"] == ind)
            n_written = sum(1 for s in f["scenarios"] if s.get("_slot", {}).get("_ind") == ind)
            cells.append(f"{n_written:>2}/{n_slots:>2}/{PLAN[code][ind][0]:>2}")
        tot = sum(v[0] for v in PLAN[code].values())
        lines.append(f"{code:<8} " + " ".join(cells) + f"   {len(f['scenarios']):>3}/{len(f['slots']):>3}/{tot}")
    return "\n".join(lines) + "\n(written / slots declared / planned)"


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--src", type=Path, default=SRC)
    ap.add_argument("--image-src", type=Path, default=IMAGE_SRC)
    ap.add_argument("--out", type=Path, default=OUT)
    ap.add_argument("--check", action="store_true", help="validate only; do not write the workbook")
    a = ap.parse_args()

    fams, sources = load(a.src)
    img_sources, img_slots = load_image_bank(a.image_src)
    rubric = load_rubric(a.image_src / "rubric.yaml")
    customers = load_customers(a.src / "customers.yaml", a.image_src / "customers.yaml")
    rep = validate(fams, sources, img_sources, img_slots, rubric, customers)
    if rep.errors:
        print(f"{len(rep.errors)} rule(s) broken:", file=sys.stderr)
        for e in rep.errors:
            print(f"  - {e}", file=sys.stderr)
        return 1
    print(progress_table(fams))
    pend = [x for x in rep.rules if x[3] in ("pending", "behind")]
    print(f"\n{len(rep.rules) - len(pend)} coverage rules met or on track, {len(pend)} pending.")
    scen = [s for f in fams.values() for s in f["scenarios"]]
    print(f"{len(scen)} scenarios, {sum(outputs_of(s) for s in scen)} clips per model, "
          f"{len(sources)} video sources, {len(img_sources)} image sources available.")
    if a.check:
        return 0
    wb = build(fams, sources, img_sources, rubric, rep, customers)
    a.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(a.out)
    print(f"wrote {a.out.relative_to(ROOT) if a.out.is_relative_to(ROOT) else a.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
