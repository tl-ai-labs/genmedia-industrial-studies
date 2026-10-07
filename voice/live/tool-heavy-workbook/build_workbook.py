"""Build the tool-heavy live-call scenario workbook (.xlsx) from src/*.yaml.

    voice/.venv/bin/python voice/live/tool-heavy-workbook/build_workbook.py

Same shape as the 216-call workbook (voice/live/scenario-workbook), for calls
that each expect 20 or more tool calls: a customer with several tasks in one
call, each task a chain of lookups and actions. Scenario ids run from 101 so
they never collide with the 216-call bank.

Every count in the workbook (partner aspects, tool calls, languages) is
derived here from the scenario sources, so the summary tabs cannot drift from
the scenarios. The script also validates each scenario and stops on errors.
"""
from __future__ import annotations

import re
import sys
from collections import Counter, OrderedDict
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
SRC = HERE / "src"
OUT = HERE / "2026-10-06-live-call-tool-heavy-scenarios-v1.xlsx"
SCENARIO_DATE = "Monday 5 October 2026, 11:00 am IST"

INDUSTRY_ORDER = ["bank", "ret", "tel", "trv", "hlt"]

LANGS = OrderedDict([
    ("en", ("English", "en")),
    ("hi", ("Hindi", "hi-hi")),
    ("hg", ("Hinglish", "hi-en")),
    ("tx", ("Telugu ↔ English (switch)", "en + te-mix")),
    ("te", ("Telugu (with English words)", "te-mix")),
    ("ta", ("Tamil (with English words)", "ta-mix")),
])
# 150 calls: English 40, Hindi 26, Hinglish 40, Telugu ↔ English 8, Telugu 18, Tamil 18.
LANG_TARGET_SHARE = {"en": 40 / 150, "hi": 26 / 150, "hg": 40 / 150, "tx": 8 / 150, "te": 18 / 150, "ta": 18 / 150}

# Planned calls per industry (the build reports progress against these).
PLANNED = {"bank": 32, "ret": 32, "tel": 27, "trv": 32, "hlt": 27}
MIN_TOOLS, MIN_LINES, MAX_LINES = 20, 12, 18

# Per-industry plan, enforced once an industry reaches its planned count.
# Languages: en, hg, hi, te, ta, tx. Tiers: 20–24, 25–29, 30+. Errors: calls with a tool error (± 2).
PLAN = {
    "bank": {"langs": (9, 9, 5, 4, 4, 1), "tiers": (16, 11, 5), "errors": 16},
    "ret":  {"langs": (9, 9, 5, 4, 4, 1), "tiers": (16, 11, 5), "errors": 16},
    "trv":  {"langs": (8, 8, 6, 4, 4, 2), "tiers": (16, 11, 5), "errors": 16},
    "tel":  {"langs": (7, 7, 5, 3, 3, 2), "tiers": (13, 9, 5), "errors": 14},
    "hlt":  {"langs": (7, 7, 5, 3, 3, 2), "tiers": (14, 8, 5), "errors": 13},
}
PLAN_LANG_ORDER = ("en", "hg", "hi", "te", "ta", "tx")

# The call happens on Monday 5 October 2026 at 11:00 am IST (SCENARIO_DATE).
CALL_DATE, CALL_MINUTES = (2026, 10, 5), 11 * 60
DOW = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
MONTHS = {"Jan": 1, "Feb": 2, "Mar": 3, "Apr": 4, "May": 5, "Jun": 6, "Jul": 7, "Aug": 8,
          "Sep": 9, "Oct": 10, "Nov": 11, "Dec": 12}
# Words that say an event has already happened; one at a later time on the call day is an error.
PAST_EVENT = r"(landed|departed|debited|credited|charged|declined|failed|delivered|dispatched|logged in|received|paid|blocked)"
# The argument that names the person or item a call acts on (for the error-retry check).
SUBJECT_ARGS = ("passenger", "claim_id", "card_last4", "account_last4", "txn_ref", "utr", "mandate_id",
                "loan_id", "fd_last4", "admission_id", "policy_no", "member", "order_id", "item_id", "awb", "return_id", "txn_id", "sku", "mobile", "fibre_id", "box_id", "complaint_id", "appointment_id", "collection_id", "rx_id", "ride_id",
                "booking_id", "patient", "slot")

ASPECTS = OrderedDict([
    ("A1", "Time to first byte"),
    ("A2", "Turn completion recognition"),
    ("A3", "Interruption recognition"),
    ("A4", "Background noise"),
    ("A5", "Response time and quality"),
    ("A6", "Tool calling"),
    ("A7", "Language switching"),
])

# Scenarios are grouped by how many tool calls they expect; the labels are shown as written.
TOOL_LEVELS = OrderedDict([
    ("20–24 calls", "Several tasks, each a short chain"),
    ("25–29 calls", "More tasks, or one item repeated for several people or orders"),
    ("30 or more calls", "The longest chains; most include a tool error to recover from"),
])

CUE_RE = re.compile(
    r"^(?P<kind>open|after|cut|bc|pause)"
    r"(?:\s+(?P<ms>\d+))?"
    r"(?:\s+(?P<ref>L\d+))?"
    r"(?:\s+wait=(?P<wait>\d+))?"
    r"(?:\s+(?P<spk>b))?$"
)


# ---------------------------------------------------------------- loading

def load_sources() -> list[dict]:
    """One header file per industry (has `industry`), plus any number of part
    files holding only `code` and `scenarios`; parts are merged in file order."""
    files = sorted(f for f in SRC.glob("*.yaml") if not f.name.startswith("_"))
    heads, parts = {}, []
    for f in files:
        data = yaml.safe_load(f.read_text(encoding="utf-8"))
        data["_file"] = f.name
        if "industry" in data:
            data.setdefault("scenarios", [])
            data["scenarios"] = list(data["scenarios"] or [])
            heads[data["code"]] = data
        else:
            parts.append(data)
    for p in parts:
        if p["code"] not in heads:
            raise SystemExit(f"{p['_file']}: no header file for industry code {p['code']}")
        heads[p["code"]]["scenarios"] += p.get("scenarios") or []
    inds = [h for h in heads.values() if h["scenarios"]]
    inds.sort(key=lambda d: INDUSTRY_ORDER.index(d["code"]))
    for ind in inds:
        for s in ind["scenarios"]:
            flatten_tasks(s)
    return inds


def flatten_tasks(s: dict) -> None:
    """`tasks: [{task, calls: [...]}]` → `tools` (in order) and `_task_of` (call index → task)."""
    if "tasks" not in s:
        s["_task_of"] = ["—"] * len(s.get("tools") or [])
        return
    s["tools"], s["_task_of"] = [], []
    for t in s["tasks"]:
        for c in t["calls"]:
            s["tools"].append(c)
            s["_task_of"].append(t["task"])


def split_call(call: str) -> tuple[str, str, str]:
    """'name(args) → result' → (name, args, result)."""
    head, _, result = call.partition("→")
    name, _, args = head.strip().partition("(")
    return name.strip(), args.rstrip().removesuffix(")").strip(), result.strip()


def is_error(call: str) -> bool:
    return split_call(call)[2].lower().startswith("error")


def parse_line(sid: str, i: int, raw) -> dict:
    if not isinstance(raw, list) or not 1 <= len(raw) <= 3:
        raise ValueError(f"{sid} line {i}: expected [cue, text, gloss?]")
    cue = raw[0].strip()
    m = CUE_RE.match(cue)
    if not m:
        raise ValueError(f"{sid} line {i}: bad cue {cue!r}")
    d = m.groupdict()
    kind = d["kind"]
    if kind in ("cut", "bc") and not (d["ms"] and d["ref"]):
        raise ValueError(f"{sid} line {i}: {kind} needs offset and line ref")
    if kind == "pause" and not d["ms"]:
        raise ValueError(f"{sid} line {i}: pause needs a gap in ms")
    text = raw[1] if len(raw) > 1 else ""
    gloss = raw[2] if len(raw) > 2 else ""
    return {
        "id": f"L{i}", "kind": kind, "ms": int(d["ms"]) if d["ms"] else None,
        "ref": d["ref"], "wait": int(d["wait"]) if d["wait"] else None,
        "speaker": "second customer" if d["spk"] else "customer",
        "text": text, "gloss": gloss,
    }


def tool_name(call: str) -> str:
    return call.split("(", 1)[0].strip()


def validate(ind: dict, s: dict, seen: set) -> list[str]:
    errs = []
    sid = s.get("id", "?")
    for k in ("id", "g", "lang", "diff", "dir", "title", "basis", "caller", "goal",
              "role", "lines", "key", "guard"):
        if k not in s:
            errs.append(f"{sid}: missing {k}")
    if sid in seen:
        errs.append(f"{sid}: duplicate id")
    seen.add(sid)
    if s.get("g") not in ind["groups"]:
        errs.append(f"{sid}: unknown group {s.get('g')}")
    if s.get("lang") not in LANGS:
        errs.append(f"{sid}: unknown lang {s.get('lang')}")
    if s.get("diff") not in ("Easy", "Medium", "Hard"):
        errs.append(f"{sid}: bad difficulty")
    if s.get("dir") not in ("in", "out"):
        errs.append(f"{sid}: bad direction")
    if s.get("lang") == "tx" and not s.get("switch"):
        errs.append(f"{sid}: tx scenario needs a switch note")
    lines = s["_lines"]
    first = lines[0]["kind"]
    if first != "open":
        errs.append(f"{sid}: every call starts with the customer on connect ('open'); "
                    "outbound calls start with the customer picking up")
    ids = [ln["id"] for ln in lines]
    for i, ln in enumerate(lines):
        if ln["kind"] == "open" and i:
            errs.append(f"{sid} {ln['id']}: 'open' only on the first line")
        if ln["kind"] in ("cut", "bc") and (i == 0 or ln["ref"] != lines[i - 1]["id"]):
            errs.append(f"{sid} {ln['id']}: must directly follow {ln['ref']}, whose reply it lands in")
        if ln["wait"] and ln["wait"] < 6000:
            errs.append(f"{sid} {ln['id']}: a silence must be at least 6 s (the agent checks in after about 5 s)")
    c = s.get("checks") or {}
    refs = list((c.get("lang_from") or {}).keys())
    refs += [c[k] for k in ("keep_talking", "no_jump_in", "reprompt_in_silence") if c.get(k)]
    if s.get("switch"):
        refs.append(s["switch"].split(":", 1)[0].strip())
    for r in refs:
        if r not in ids:
            errs.append(f"{sid}: check or switch refers to {r}, which is not a line")
    facts = re.sub(r"[\s,]", "", " ".join([s.get("context") or ""] + list(s.get("tools") or [])))
    for d in c.get("digits") or []:
        if re.sub(r"[\s,]", "", str(d)) not in facts:
            errs.append(f"{sid}: digits_exact {d} is not in the facts or tool results")
        if ln["kind"] in ("cut", "bc", "pause") and not ln["text"].strip():
            errs.append(f"{sid} {ln['id']}: empty text")
    for call in s.get("tools") or []:
        name = tool_name(call)
        if name not in ind["tools"]:
            errs.append(f"{sid}: tool {name} not in the {ind['code']} catalogue")
    for f in s.get("forbid") or []:
        if tool_name(f) not in ind["tools"]:
            errs.append(f"{sid}: forbidden tool {f} not in the catalogue")
        if "(" in f and any(call.split("→")[0].strip() == f.strip() for call in s.get("tools") or []):
            errs.append(f"{sid}: forbidden call {f} is also an expected call")
    # tool-heavy rules
    if not re.fullmatch(rf"{ind['code']}-1\d\d", sid):
        errs.append(f"{sid}: ids are <code>-101 … <code>-199 in this bank")
    n_tools = len(s.get("tools") or [])
    if n_tools < MIN_TOOLS:
        errs.append(f"{sid}: {n_tools} tool calls; this bank needs {MIN_TOOLS} or more")
    if not MIN_LINES <= len(lines) <= MAX_LINES:
        errs.append(f"{sid}: {len(lines)} customer lines; this bank uses {MIN_LINES}–{MAX_LINES}")
    if not s.get("forbid"):
        errs.append(f"{sid}: every call in this bank has a call the agent must not make")
    names = {tool_name(t) for t in s.get("tools") or []}
    for o in c.get("tool_order") or []:
        if tool_name(o) not in names:
            errs.append(f"{sid}: tool_order names {o}, which is not an expected call")
    for call in s.get("tools") or []:
        if "→" not in call:
            errs.append(f"{sid}: call without a fixed result: {call}")
    return errs


def _minutes(t: str) -> int:
    h, m, ap = re.match(r"(\d{1,2}):(\d{2}) (am|pm)", t).groups()
    return int(h) % 12 * 60 + (720 if ap == "pm" else 0) + int(m)


def validate_facts(s: dict) -> list[str]:
    """Weekdays match the 2026 calendar; every tool error is retried; nothing on the
    call day has already happened after the 11:00 am call time."""
    import datetime
    errs, sid = [], s["id"]
    texts = list(s.get("tools") or []) + [s.get("context") or "", s.get("role") or ""] + \
        [str(x) for raw in s["lines"] for x in raw[1:]] + list(s.get("key") or [])
    for t in texts:
        for d, n, m in re.findall(r"\b(Mon|Tue|Wed|Thu|Fri|Sat|Sun)\w* (\d{1,2}) (Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)", t):
            # Up to six months before the call is the past (2026); earlier months are next year.
            year = 2026 if MONTHS[m] >= CALL_DATE[1] - 6 else 2027
            real = DOW[datetime.date(year, MONTHS[m], int(n)).weekday()]
            if real != d:
                errs.append(f"{sid}: '{d} {n} {m}' — {n} {m} {year} is a {real}")
    calls = [split_call(c) for c in s.get("tools") or []]
    for i, (name, args, res) in enumerate(calls):
        if not res.lower().startswith("error"):
            continue
        subj = next((m.group(0) for a in SUBJECT_ARGS for m in [re.search(rf"\b{a}=[^,]+", args)] if m), None)
        if not any(c[0] == name and (subj is None or subj in c[1]) for c in calls[i + 1:]):
            errs.append(f"{sid}: {name}({args}) returns an error and is never retried")
    for name, args, res in calls:
        if not re.search(r"\bMon 5 Oct\b|\btoday\b", args + " " + res):
            continue
        for verb, t in re.findall(PAST_EVENT + r"\s+(?:at\s+)?(\d{1,2}:\d{2} (?:am|pm))", res):
            if _minutes(t) > CALL_MINUTES:
                errs.append(f"{sid}: '{verb} {t}' on the call day is after the 11:00 am call")
    return errs


def check_plan(inds: list[dict]) -> tuple[list[str], list[str]]:
    """(errors, notes): an industry at its planned count must match its plan."""
    errs, notes = [], []
    for ind in inds:
        code, ss = ind["code"], ind["scenarios"]
        if len(ss) < PLANNED[code]:
            notes.append(f"{code}: {len(ss)}/{PLANNED[code]} built — plan not checked yet")
            continue
        if len(ss) > PLANNED[code]:
            errs.append(f"{code}: {len(ss)} calls, plan {PLANNED[code]}")
        plan = PLAN[code]
        lc = Counter(s["lang"] for s in ss)
        for k, want in zip(PLAN_LANG_ORDER, plan["langs"]):
            if lc[k] != want:
                errs.append(f"{code}: {LANGS[k][0]} {lc[k]}, plan {want}")
        tc = Counter(tool_level(s) for s in ss)
        for k, want in zip(TOOL_LEVELS, plan["tiers"]):
            if tc[k] != want:
                errs.append(f"{code}: {k} {tc[k]}, plan {want}")
        n_err = sum(1 for s in ss if any(is_error(c) for c in s["tools"]))
        if abs(n_err - plan["errors"]) > 2:
            errs.append(f"{code}: {n_err} calls with a tool error, plan {plan['errors']} (± 2)")
    return errs, notes


# ---------------------------------------------------------------- derived fields

def tool_level(s: dict) -> str:
    n = len(s.get("tools") or [])
    return "20–24 calls" if n < 25 else "25–29 calls" if n < 30 else "30 or more calls"


def aspects(s: dict) -> list[str]:
    kinds = {ln["kind"] for ln in s["_lines"]}
    waits = any(ln["wait"] for ln in s["_lines"])
    out = ["A1"]
    if "pause" in kinds or waits:
        out.append("A2")
    if kinds & {"cut", "bc"}:
        out.append("A3")
    if s.get("noise"):
        out.append("A4")
    out.append("A5")
    if s.get("tools"):
        out.append("A6")
    if s.get("switch"):
        out.append("A7")
    return out


def cue_plain(ln: dict) -> str:
    k = ln["kind"]
    who = " (second customer)" if ln["speaker"] != "customer" else ""
    if k == "open":
        return "When the call connects (customer picks up / speaks first)" + who
    if k == "after":
        base = "After the agent finishes (1.0 s of agent silence)"
        if ln["wait"]:
            base = f"Customer stays silent {ln['wait'] / 1000:g} s after the agent finishes, then speaks"
        return base + who
    if k == "cut":
        return f"INTERRUPTION — {ln['ms'] / 1000:g} s into the agent's reply to {ln['ref']}" + who
    if k == "bc":
        return f"ACKNOWLEDGEMENT (agent should keep talking) — {ln['ms'] / 1000:g} s into the reply to {ln['ref']}" + who
    if k == "pause":
        return f"PAUSE — continues {ln['ms'] / 1000:g} s after the previous line ends; agent should wait" + who
    return k


def planted(s: dict) -> str:
    out = []
    for ln in s["_lines"]:
        if ln["kind"] == "cut":
            out.append(f"{ln['id']}: interruption {ln['ms']} ms into the reply to {ln['ref']}")
        elif ln["kind"] == "bc":
            out.append(f"{ln['id']}: acknowledgement {ln['ms']} ms into the reply to {ln['ref']} (keep talking)")
        elif ln["kind"] == "pause":
            out.append(f"{ln['id']}: mid-sentence pause of {ln['ms']} ms (wait)")
        elif ln["wait"]:
            out.append(f"{ln['id']}: customer silent {ln['wait']} ms before speaking (check in once)")
        if ln["speaker"] != "customer" and (not out or ln["id"] not in out[-1]):
            pass
    speakers = [ln for ln in s["_lines"] if ln["speaker"] != "customer"]
    if speakers:
        out.append(f"{speakers[0]['id']}: phone handed to a second customer")
    if s.get("switch"):
        out.append("Language switch — " + s["switch"])
    if s.get("noise"):
        out.append("Background sound — " + noise_text(s["noise"]))
    return "\n".join(out) or "—"


def noise_text(n: dict | None) -> str:
    if not n:
        return "—"
    t = f"{n['bed'].replace('_', ' ')}, {n['snr_db']} dB below the customer, whole call"
    if n.get("decoy"):
        t += f"; decoy voice: {n['decoy']}"
    return t


def script_text(s: dict, with_gloss: bool = True) -> str:
    rows = []
    for ln in s["_lines"]:
        tag = {"open": "", "after": "", "cut": " [INTERRUPTS]", "bc": " [ACK]", "pause": " [after pause]"}[ln["kind"]]
        if ln["wait"]:
            tag = f" [after {ln['wait'] / 1000:g} s silence]"
        who = " (2nd)" if ln["speaker"] != "customer" else ""
        row = f"{ln['id']}{who}{tag}: {ln['text']}"
        if with_gloss and ln["gloss"]:
            row += f"\n      — {ln['gloss']}"
        rows.append(row)
    return "\n".join(rows)


def profile_text(s: dict) -> str:
    return s["caller"].replace("then speaker b:", "then a second person:").replace("speaker b:", "then a second person:")


def agent_told(s: dict) -> str:
    """The scenario's own part of the agent instructions: role, then facts and policy."""
    parts = [s["role"].strip()]
    if (s.get("context") or "").strip():
        parts.append(s["context"].strip())
    return "\n\n".join(parts)


def tools_text(s: dict) -> str:
    calls = s.get("tools") or []
    n = len(calls)
    out = [f"{i}. {t}" for i, t in enumerate(calls, 1)] or ["No tools — the facts are in the instructions"]
    order = (s.get("checks") or {}).get("tool_order")
    if order:
        out.append("Order matters: " + " before ".join(order))
    if s.get("forbid"):
        out.append("Must not call (fails the call): " + "; ".join(s["forbid"]))
    return "\n".join(out)


LANG_NAME = {"en": "English", "hi": "Hindi", "hg": "Hinglish", "te": "Telugu", "ta": "Tamil", "hi-en": "Hinglish"}
ARG_LABEL = {"card_last4": "card ending", "account_last4": "account ending", "pincode": "pincode",
             "mobile": "mobile number", "fd_last4": "FD ending", "utr": "UTR", "cheque": "cheque",
             "Flat": "flat"}


def number_label(digits: str, s: dict) -> str:
    """Show a number the way it appears in the facts, e.g. 'DSP 64420' or 'card ending 5821'."""
    facts = " ".join([s.get("context") or ""] + list(s.get("tools") or []))
    pat = r"[\s,]?".join(re.escape(ch) for ch in re.sub(r"[\s,]", "", digits))
    m = re.search(pat, facts)
    if not m:
        return digits
    before, after = facts[:m.start()], facts[m.end():]
    unit = re.match(r"\s*(months|days)\b", after)
    if unit:
        return f"{digits} {unit.group(1)}"
    word = re.search(r"([A-Za-z_][A-Za-z_0-9]*)\W{0,3}$", before)
    if not word:
        return digits
    w = word.group(1)
    if re.fullmatch(r"[1-9]\d{5}", re.sub(r"\s", "", digits)) and w not in ARG_LABEL:
        return f"pincode {digits}"
    if w in ARG_LABEL:
        return f"{ARG_LABEL[w]} {digits}"
    if w.isupper() and 2 <= len(w) <= 4:
        return f"{w} {digits}"
    return digits


def moments_text(s: dict) -> str:
    from sheets_text import NOISE_BEDS
    """Each planted moment with what the agent must do, and the rubric aspect it feeds."""
    c = s.get("checks") or {}
    out = []
    for ln in s["_lines"]:
        k, i = ln["kind"], ln["id"]
        if k == "cut":
            out.append(f"{i} · customer cuts in {ln['ms'] / 1000:g} s into the agent's reply → agent stops "
                       f"within 1 s and answers the cut-in (A3)")
        elif k == "bc":
            out.append(f"{i} · customer says a short '{ln['text'].strip(' .…!')}' while the agent talks → agent "
                       f"keeps talking (A3)")
        elif k == "pause":
            out.append(f"{i} · customer stops mid-sentence for {ln['ms'] / 1000:g} s → agent waits for the rest (A2)")
        elif ln["wait"]:
            out.append(f"{i} · customer is silent for {ln['wait'] / 1000:g} s → agent checks in once (A2)")
    second = [ln for ln in s["_lines"] if ln["speaker"] != "customer"]
    if second:
        out.append(f"{second[0]['id']} · a second person takes the phone → agent carries on with them")
    for line, lg in (c.get("lang_from") or {}).items():
        frm = s["switch"].split(":", 1)[1].split("→")[0].strip() if s.get("switch") else ""
        frm = re.sub(r"\s*\(.*", "", frm)
        out.append(f"{line} · language changes{' from ' + frm if frm else ''} to {LANG_NAME.get(lg, lg)} → "
                   f"agent replies in {LANG_NAME.get(lg, lg)} from its next reply (A7)")
    n = s.get("noise")
    if n:
        name = next(b["what"] for b in NOISE_BEDS if b["bed"] == n["bed"]).split(" (")[0]
        out.append(f"Whole call · background sound: {name if name[1:2].isupper() else name[0].lower() + name[1:]}, {n['snr_db']} dB quieter than "
                   f"the customer → task still done (A4)")
        if n.get("decoy"):
            out.append(f"Whole call · decoy: {n['decoy']} → agent never repeats its numbers (A4)")
    return "\n".join(out) or "None — a plain call with no planted moments"


def moments_plain(s: dict) -> str:
    """What happens in the call, by line — no expectations (those are in the check columns)."""
    out = []
    for ln in s["_lines"]:
        k, i = ln["kind"], ln["id"]
        if k == "cut":
            out.append(f"{i} · customer cuts in {ln['ms'] / 1000:g} s into the agent's reply")
        elif k == "bc":
            out.append(f"{i} · customer says a short '{ln['text'].strip(' .…!')}' while the agent is talking")
        elif k == "pause":
            out.append(f"{i} · customer stops mid-sentence for {ln['ms'] / 1000:g} s, then carries on")
        elif ln["wait"]:
            out.append(f"{i} · customer stays silent for {ln['wait'] / 1000:g} s before speaking")
    second = [ln for ln in s["_lines"] if ln["speaker"] != "customer"]
    if second:
        out.append(f"{second[0]['id']} · a second person takes the phone")
    if s.get("switch"):
        line, rest = s["switch"].split(":", 1)
        rest = rest.replace("speaker b", "second person").replace(", second person", "")
        out.append(f"{line.strip()} · language changes: {rest.strip()}")
    n = s.get("noise")
    if n:
        name = next(b_["what"] for b_ in _beds() if b_["bed"] == n["bed"]).split(" (")[0]
        out.append(f"Whole call · background sound: {name[0].lower() + name[1:] if not name[1:2].isupper() else name}"
                   + (f" with a decoy: {n['decoy']}" if n.get("decoy") else ""))
    return "\n".join(out) or "None — a plain call"


def automated_checks(s: dict) -> str:
    """Everything code checks for this call, each line led by its rubric code."""
    c = s.get("checks") or {}
    rows = [("A1", "A1 · A5 · how fast the agent starts and answers, on every turn")]
    for ln in s["_lines"]:
        i = ln["id"]
        if ln["kind"] == "pause":
            rows.append(("A2", f"A2 · {i}: waits through the customer's {ln['ms'] / 1000:g} s pause"))
        if ln["wait"]:
            rows.append(("A2", f"A2 · {i}: checks in once during the {ln['wait'] / 1000:g} s silence"))
        if ln["kind"] == "cut":
            rows.append(("A3", f"A3 · {i}: stops talking within 1 s when the customer cuts in"))
        if ln["kind"] == "bc":
            rows.append(("A3", f"A3 · {i}: keeps talking through '{ln['text'].strip(' .…!')}'"))
    n = s.get("noise")
    decoy_words = []
    if n:
        decoy = n.get("decoy") or ""
        decoy_words = [x for x in c.get("not_say") or [] if re.sub(r"\D", "", x) and re.sub(r"\D", "", x) in re.sub(r"\D", "", decoy + " ".join(b.get("decoy", "") for b in _beds()))]
        rows.append(("A4", "A4 · gets the task and numbers right despite the background sound"
                     + (" and never repeats its numbers" if n.get("decoy") else "")))
    if s.get("tools"):
        order = c.get("tool_order")
        rows.append(("A6", "A6 · makes the tool calls in the Tools column, with the same details"
                     + (" — " + " before ".join(order) if order else "")))
    for line, lg in (c.get("lang_from") or {}).items():
        rows.append(("A7", f"A7 · {line}: replies in {LANG_NAME.get(lg, lg)} from its next reply"))
    if s.get("forbid"):
        rows.append(("P5", "P5 · never calls " + " or ".join(s["forbid"])))
    say = []
    if c.get("digits"):
        say.append(" · ".join(number_label(str(d), s) for d in c["digits"]))
    for alts in c.get("say") or []:
        say.append(" or ".join(f"'{a}'" for a in alts))
    if say:
        rows.append(("P9", "P9 · says: " + "; ".join(say)))
    banned = [x for x in c.get("not_say") or [] if x not in decoy_words]
    otp = [x for x in banned if re.search(r"OTP|PIN|CVV", x)]
    if otp:
        rows.append(("P3", "P3 · never says " + " / ".join(f"'{x}'" for x in otp)))
    other = [x for x in banned if x not in otp] + decoy_words
    if other:
        rows.append(("P9", "P9 · never says " + " / ".join(f"'{x}'" for x in other)))
    if not c.get("lang_from"):
        rows.append(("P10", f"P10 · replies in {LANG_NAME.get(c.get('lang') or s['lang'], s['lang'])} throughout"))
    order = ["A1", "A2", "A3", "A4", "A6", "A7", "P3", "P5", "P9", "P10"]
    rows.sort(key=lambda r: order.index(r[0]))
    return "\n".join(r[1] for r in rows)


def _beds():
    from sheets_text import NOISE_BEDS
    return NOISE_BEDS


def judge_checks(s: dict) -> str:
    """The answer key, filed under the judge's rubric codes (code-checked items removed)."""
    from checks_split import judge_items
    order = ["J1", "J2", "J3", "J4", "J5", "J6", "P3", "P4", "P6", "P7", "P8"]
    items = sorted(judge_items(s["id"], s["key"]), key=lambda x: order.index(x[0]))
    return "\n".join(f"{code} · {text}" for code, text in items) or "— (everything in this call is checked by code)"


def words_text(s: dict) -> str:
    """The exact words this call listens for, in plain language."""
    c = s.get("checks") or {}
    out = []
    if not c.get("lang_from"):
        out.append(f"Replies in {LANG_NAME.get(c.get('lang') or s['lang'], s['lang'])} throughout")
    if c.get("digits"):
        out.append("Says exactly: " + " · ".join(number_label(str(d), s) for d in c["digits"]))
    for alts in c.get("say") or []:
        out.append(("Says: " if len(alts) == 1 else "Says at least one of: ") + " / ".join(str(a) for a in alts))
    if c.get("not_say"):
        out.append("Never says: " + " / ".join(f"'{x}'" for x in c["not_say"]))
    return "\n".join(out) or "— (no exact words for this call; the language is checked under Planted moments)"


def checks_text(s: dict) -> str:
    c = s.get("checks") or {}
    out = ["time_to_first_byte: every turn, from our recording (median and 95th percentile)"]
    lang = s["lang"]
    if lang == "tx" or c.get("lang_from"):
        lf = c.get("lang_from") or {}
        for line, lg in lf.items():
            out.append(f"reply_language: {lg} from {line} onward (switch within one reply)")
    else:
        out.append(f"reply_language: {c.get('lang', {'en': 'en', 'hi': 'hi', 'hg': 'hi-en', 'te': 'te', 'ta': 'ta'}.get(lang, lang))} every turn")
    for ln in s["_lines"]:
        if ln["kind"] == "cut":
            out.append(f"stop_after_interruption: agent audio stops ≤ 1000 ms after {ln['id']} starts")
        if ln["kind"] == "bc":
            out.append(f"keep_talking: agent does not stop for {ln['id']} (an acknowledgement)")
        if ln["kind"] == "pause":
            out.append(f"no_jump_in: agent does not start a reply in the pause before {ln['id']}")
        if ln["wait"]:
            out.append(f"check_in_on_silence: agent checks in once during the silence before {ln['id']}")
    for d in c.get("digits") or []:
        out.append(f"digits_exact: {d}")
    for alts in c.get("say") or []:
        out.append("must_say (any of): " + " | ".join(str(a) for a in alts))
    for ns in c.get("not_say") or []:
        out.append(f"must_not_say: {ns}")
    if s.get("tools"):
        out.append(f"tool_calls: {len(s['tools'])} expected, names and key arguments exact")
    if c.get("tool_order"):
        out.append("tool_order: " + " before ".join(c["tool_order"]))
    for f in s.get("forbid") or []:
        out.append(f"tool_not_called: {f}")
    if s.get("noise", {}) and s["noise"].get("decoy"):
        out.append("decoy_ignored: no digits from the decoy voice in any reply")
    return "\n".join(out)


def pass_rule(s: dict) -> str:
    return ("Pass = every answer-key item met, every automatic check passes, and no failed "
            "completion or safety check. A content check (digits, names) that fails costs points "
            "and is shown on the call.")


# ---------------------------------------------------------------- workbook helpers

HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
SECTION_FONT = Font(bold=True, size=11, color="1F3A5F")
WRAP = Alignment(wrap_text=True, vertical="top")


def table(ws, headers, rows, widths, freeze="B2"):
    ws.append(headers)
    for c in ws[1]:
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
        c.alignment = WRAP
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions


def text_sheet(ws, title, blocks, width=120):
    ws.column_dimensions["A"].width = width
    ws.append([title])
    ws["A1"].font = TITLE_FONT
    for heading, body in blocks:
        ws.append([])
        ws.append([heading])
        ws.cell(ws.max_row, 1).font = SECTION_FONT
        for para in body:
            ws.append([para])
            ws.cell(ws.max_row, 1).alignment = WRAP


# ---------------------------------------------------------------- sheets

def build(inds: list[dict]) -> Workbook:
    from sheets_text import (DECOY_TTS, ALSO_RELEVANT, AGENT_BASE, READ_ME_COLUMNS, READ_ME_SUMMARY,
                             NOISE_BEDS, PROFILES)

    all_s = [s for ind in inds for s in ind["scenarios"]]
    wb = Workbook()

    # 1 · Method
    ws = wb.active
    ws.title = "Read me"
    ws.append(["Live-Call Scenarios — tool-heavy calls (20+ tool calls each), version 1"])
    ws["A1"].font = TITLE_FONT
    for para in READ_ME_SUMMARY(all_s, inds):
        ws.append([para])
        ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=3)
        ws.cell(ws.max_row, 1).alignment = WRAP
        ws.row_dimensions[ws.max_row].height = 48
    ws.append([])
    ws.append(["Tab", "Column", "What it means"])
    for c in ws[ws.max_row]:
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, WRAP
    for tab, col, meaning in READ_ME_COLUMNS:
        ws.append([tab, col or "(the tab)", meaning])
        for c in ws[ws.max_row]:
            c.alignment = WRAP
        if not col:
            ws.cell(ws.max_row, 1).font = Font(bold=True)
    for i, w in enumerate([24, 34, 110], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # 2 · Scenario Bank
    ws = wb.create_sheet("Scenario Bank")
    headers = ["Scenario ID", "Industry", "Drill-down group", "Language", "Difficulty",
               "Direction", "What this call tests", "Scenario", "Real-world basis",
               "Customer", "Customer goal", "What the agent is told for this call",
               "Tools (expected calls → fixed results)",
               "Background sound", "Planted moments (what happens in the call)",
               "Conversation script (customer lines, in order)", "Automated checks (by code)",
               "Judge checks (by the AI judge)",
               "Native-speaker review needed",
               "Tool calls (number)", "Tasks in the call", "Tool errors the agent must handle"]
    rows = []
    for ind in inds:
        for s in ind["scenarios"]:
            lname, script = LANGS[s["lang"]]
            rows.append([
                s["id"], ind["industry"], s["g"], lname, s["diff"],
                "Inbound — customer speaks first" if s["dir"] == "in"
                else "Agent opens after the customer's hello",
                ", ".join(f"{a} {ASPECTS[a]}" for a in aspects(s)),
                s["title"], s["basis"], profile_text(s),
                s["goal"], agent_told(s),
                tools_text(s),
                noise_text(s.get("noise")), moments_plain(s),
                script_text(s), automated_checks(s), judge_checks(s),
                "Yes" if s["lang"] in ("te", "ta", "tx") else "No",
                len(s["tools"]),
                "\n".join(f"{i}. {t}" for i, t in enumerate(dict.fromkeys(s["_task_of"]), 1)),
                "\n".join(f"{i}. {split_call(c)[0]} → {split_call(c)[2]}"
                          for i, c in enumerate(s["tools"], 1) if is_error(c)) or "—",
            ])
    table(ws, headers, rows,
          [12, 18, 10, 16, 9, 13, 30, 34, 40, 24, 30, 60, 80, 30, 40, 90,
           55, 55, 10, 10, 40, 50])

    # 2b · Tool calls: one row per expected call
    ws = wb.create_sheet("Tool calls")
    rows = []
    for ind in inds:
        for s in ind["scenarios"]:
            order = (s.get("checks") or {}).get("tool_order") or []
            for i, (call, task) in enumerate(zip(s["tools"], s["_task_of"]), 1):
                name, args, result = split_call(call)
                rule = ""
                if name in order and order.index(name):
                    rule = "after " + order[order.index(name) - 1]
                rows.append([s["id"], i, task, name, args, result,
                             "Yes — the agent must recover" if is_error(call) else "No", rule])
            for f in s.get("forbid") or []:
                rows.append([s["id"], "—", "Must not call", tool_name(f), f.partition("(")[2].rstrip(")"),
                             "Making this call fails the call (P5)", "—", ""])
    table(ws, ["Scenario ID", "Step", "Task (what the customer asked for)", "Tool", "Arguments",
               "Fixed result", "Error result?", "Order rule"],
          rows, [12, 6, 40, 26, 60, 70, 16, 24], freeze="C2")

    # 3 · Industry mapping
    ws = wb.create_sheet("Industry mapping")
    rows = []
    for ind in inds:
        for s in ind["scenarios"]:
            rows.append([s["id"], "Voice — live call", LANGS[s["lang"]][0],
                         f"{ind['industry']} — {ind['groups'][s['g']]['stage'].split(' — ')[0]}",
                         s["title"], s["basis"], ind["industry"], ALSO_RELEVANT[s["g"]]])
    table(ws, ["Scenario ID", "Modality", "Language", "Use case family", "Scenario", "Real-world basis",
               "Primary industry", "Also relevant to"], rows, [12, 16, 22, 42, 44, 60, 28, 44], freeze="A2")

    # 4 · Customer lines (the recording script)
    ws = wb.create_sheet("Customer lines")
    rows = []
    for ind in inds:
        for s in ind["scenarios"]:
            for ln in s["_lines"]:
                prof = profile_text(s)
                rows.append([s["id"], ln["id"], ln["speaker"], prof, LANGS[s["lang"]][0],
                             cue_plain(ln), ln["text"], ln["gloss"] or "(same — English)",
                             noise_text(s.get("noise")) if s.get("noise") else "—"])
    table(ws, ["Scenario ID", "Line", "Speaker", "Customer profile", "Language", "When it plays",
               "What the input voice says (verbatim)", "English meaning", "Background sound under it"],
          rows, [12, 6, 14, 30, 18, 42, 70, 60, 34], freeze="C2")

    # 4b · Input voice prompts (how each customer line is generated)
    ws = wb.create_sheet("Input voice prompts")
    rows = input_voice_rows(inds)
    table(ws, ["Scenario ID", "Line", "Speaker", "Gender · age", "Voice pool (one drawn per speaker, by seed)",
               "Language code", "Style (sent as speech_metadata.style — not spoken)",
               "Text to speak (sent verbatim)", "English meaning", "How it is played in the call",
               "Model", "After generation"],
          rows, [12, 6, 16, 10, 40, 16, 60, 70, 55, 40, 22, 50], freeze="C2")

    # 5 · Rubrics: three short tables in the same simple shape
    from sheets_text import (RUBRIC_INTRO, RUBRIC_STEP1, RUBRIC_STEP2, RUBRIC_STEP2_NOTE, RUBRIC_STEP3,
                             RUBRIC_STEP3_NOTE)
    ws = wb.create_sheet("Rubrics")
    for i_, w in enumerate([7, 30, 54, 28, 28, 28, 28], 1):
        ws.column_dimensions[get_column_letter(i_)].width = w
    ws.append(["How each call is scored"])
    ws["A1"].font = TITLE_FONT
    for para in RUBRIC_INTRO:
        ws.append([para])
        ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=7)

    def step(title, headers, rows, note=None):
        ws.append([])
        ws.append([title])
        ws.cell(ws.max_row, 1).font = SECTION_FONT
        ws.append(headers)
        for c in ws[ws.max_row]:
            c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, WRAP
        for r in rows:
            ws.append(r)
            for c in ws[ws.max_row]:
                c.alignment = WRAP
        if note:
            ws.append([note])
            ws.cell(ws.max_row, 1).font = Font(italic=True)
            ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=7)

    step("Step 1 · Pass-fail checks — does the call count?",
         ["Code", "Check", "Checked by", "If it fails"], RUBRIC_STEP1)
    step("Step 2 · Automated checks — measured by code",
         ["Code", "Aspect", "What is measured", "Good", "Acceptable", "Poor"], RUBRIC_STEP2, RUBRIC_STEP2_NOTE)
    step("Step 3 · Judge checks — the whole conversation",
         ["Code", "Criterion", "What the judge asks", "Weight", "5 — best", "3", "1 — worst"], RUBRIC_STEP3,
         RUBRIC_STEP3_NOTE)

    # 7 · Use-case profiles
    ws = wb.create_sheet("Use case profiles")
    table(ws, ["Industry", "Weight the heavy ones", "Extra fail rule for this industry", "Brand used (fictional)"],
          PROFILES, [26, 60, 70, 30])

    # 8 · Tool levels and catalogue
    ws = wb.create_sheet("Tools")
    rows = []
    for ind in inds:
        used = Counter(tool_name(t) for s in ind["scenarios"] for t in s.get("tools") or [])
        for name, desc in ind["tools"].items():
            rows.append([ind["industry"], name, desc, used[name]])
    table(ws, ["Industry", "Tool", "Arguments → returns", "Expected calls across scenarios"],
          rows, [24, 26, 80, 14])
    ws.append([])
    ws.append(["Scenarios by number of tool calls"])
    ws.cell(ws.max_row, 1).font = SECTION_FONT
    lvl = Counter(tool_level(s) for s in all_s)
    for k, v in TOOL_LEVELS.items():
        ws.append([k, v, "", lvl[k]])
    ws.append(["Tool calls in total", "", "", sum(len(s["tools"]) for s in all_s)])
    ws.append(["… of which return an error", "", "", sum(is_error(c) for s in all_s for c in s["tools"])])

    # 9 · Background sounds: prompts to make each one
    ws = wb.create_sheet("Background sounds")
    used = {}
    for s_ in all_s:
        if s_.get("noise"):
            used.setdefault(s_["noise"]["bed"], []).append(s_["id"])
    known = {n["bed"] for n in NOISE_BEDS}
    if set(used) - known:
        raise SystemExit(f"background sounds without a prompt: {sorted(set(used) - known)}")
    rows = [[n["bed"].replace("_", " "), n["what"], n["prompt"],
             n.get("decoy", "— (no decoy)"), n.get("decoy_en", "—"),
             f"{n['voice']} · {DECOY_TTS} · style: {n['style']}" if n.get("decoy") else "—",
             ", ".join(used[n["bed"]])]
            for n in NOISE_BEDS if n["bed"] in used]
    table(ws, ["Background sound", "What it is", "Sound prompt (ambience)", "Decoy voice — text (verbatim)",
               "Decoy voice — English meaning", "Decoy voice — voice, model and style", "Used in"],
          rows, [18, 30, 70, 60, 50, 40, 28])

    # 10 · Agent base instructions
    ws = wb.create_sheet("Agent base instructions")
    rows = [["All industries", AGENT_BASE.strip()]]
    for ind in inds:
        rows.append([ind["industry"], ind["agent_role_base"].strip()])
    table(ws, ["Applies to", "Instructions (sent identically to both systems)"], rows, [26, 120])

    # 11 · Drill-down
    ws = wb.create_sheet("Drill-down")
    rows = []
    for ind in inds:
        for g, meta in ind["groups"].items():
            ss = [s for s in ind["scenarios"] if s["g"] == g]
            lc = Counter(s["lang"] for s in ss)
            tl = Counter(tool_level(s) for s in ss)
            ac = Counter(a for s in ss for a in aspects(s))
            rows.append([ind["industry"], g, meta["stage"], meta["accounts"], meta["tests"],
                         ", ".join(s["id"] for s in ss), len(ss),
                         *[lc[k] for k in LANGS], *[ac[a] for a in ("A2", "A3", "A4", "A6", "A7")],
                         *[tl[k] for k in TOOL_LEVELS]])
    table(ws, ["Industry", "Drill-down group", "Journey stage", "Example accounts (relevance only)", "What it tests",
               "Scenario IDs", "Total", *[LANGS[k][0] for k in LANGS],
               "Turn completion", "Interruption", "Noise", "Tool calling", "Language switching",
               *TOOL_LEVELS.keys()],
          rows, [22, 8, 40, 34, 40, 50, 7] + [9] * 6 + [9] * 5 + [6] * 5, freeze="C2")

    # 12 · Summary tables
    ws = wb.create_sheet("Summary tables")
    summary(ws, inds, all_s)

    return wb


SCRIPT_RE = {"hi-IN": re.compile(r"[\u0900-\u097F]"), "te-IN": re.compile(r"[\u0C00-\u0C7F]"),
             "ta-IN": re.compile(r"[\u0B80-\u0BFF]")}
PROFILE_RE = re.compile(r"([MF]) (13-17|18-25|26-45|46-60|60\+)")


def speaker_profile(s, second):
    text = s["caller"]
    if second:
        if "speaker b" not in text:
            raise SystemExit(f"{s['id']}: second-customer lines but no 'speaker b' in the profile")
        text = text.split("speaker b", 1)[1]
    else:
        text = text.split("speaker b", 1)[0]
    m = PROFILE_RE.search(text)
    if not m:
        raise SystemExit(f"{s['id']}: cannot read gender and age from {text!r}")
    return m.group(1), m.group(2), text[m.end():]


def line_language(s, ln, idx):
    for code, rx in SCRIPT_RE.items():
        if rx.search(ln["text"]):
            return code
    default = {"en": "en-IN", "hi": "hi-IN", "hg": "hi-IN romanised", "te": "te-IN", "ta": "ta-IN", "tx": "en-IN"}
    switch_at = int(s["switch"].split(":")[0].strip()[1:]) if s.get("switch") else None
    if switch_at and idx >= switch_at and s["lang"] in ("te", "ta", "hg", "hi", "en"):
        return "en-IN" if s["lang"] in ("te", "ta", "hg") else default[s["lang"]]
    return default[s["lang"]]


def input_voice_rows(inds):
    from sheets_text import (AGE_STYLE, CUE_STYLE, MOOD_WORDS, PROCESSING, STYLE_BASE, TTS_MODEL,
                             VOICE_POOLS)
    rows = []
    for ind in inds:
        for s in ind["scenarios"]:
            lines = s["_lines"]
            for i, ln in enumerate(lines):
                second = ln["speaker"] != "customer"
                g, age, rest = speaker_profile(s, second)
                pool = VOICE_POOLS[(g, age)]
                if second:
                    first_pool = VOICE_POOLS[speaker_profile(s, False)[:2]]
                    pool = [v for v in pool if v not in first_pool] or pool
                lang = line_language(s, ln, i + 1)
                style = [STYLE_BASE[lang]]
                if age in AGE_STYLE:
                    style.append(AGE_STYLE[age])
                low = rest.lower()
                style += [w for w in MOOD_WORDS if re.search(rf"\b{re.escape(w)}\b", low)]
                if ln["kind"] in CUE_STYLE:
                    style.append(CUE_STYLE[ln["kind"]])
                if i + 1 < len(lines) and lines[i + 1]["kind"] == "pause":
                    style.append(CUE_STYLE["before_pause"])
                if ln["wait"]:
                    style.append(CUE_STYLE["wait"])
                rows.append([s["id"], ln["id"], ln["speaker"] + (" (2nd voice)" if second else ""),
                             f"{g} {age}", ", ".join(pool), lang, ", ".join(dict.fromkeys(style)),
                             ln["text"], ln["gloss"] or "(same — English)", cue_plain(ln),
                             TTS_MODEL, PROCESSING])
    return rows


def summary(ws, inds, all_s):
    def block(title, headers, rows):
        ws.append([title])
        ws.cell(ws.max_row, 1).font = SECTION_FONT
        ws.append(headers)
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append(r)
        ws.append([])

    n = len(all_s)
    lc = Counter(s["lang"] for s in all_s)
    block("1. Scenarios by language", ["Language", "Scenarios", "Share", "Target share"],
          [[LANGS[k][0], lc[k], f"{lc[k] / n:.1%}", f"{LANG_TARGET_SHARE[k]:.1%}"] for k in LANGS]
          + [["All", n, "100%", ""]])
    block("2. Scenarios by industry and language", ["Industry", *[LANGS[k][0] for k in LANGS], "Total"],
          [[ind["industry"], *[sum(1 for s in ind["scenarios"] if s["lang"] == k) for k in LANGS],
            len(ind["scenarios"])] for ind in inds])
    block("3. Partner aspects planted, by industry",
          ["Industry", *[f"{k} {v}" for k, v in ASPECTS.items()]],
          [[ind["industry"], *[sum(1 for s in ind["scenarios"] if a in aspects(s)) for a in ASPECTS]]
           for ind in inds]
          + [["All", *[sum(1 for s in all_s if a in aspects(s)) for a in ASPECTS]]])
    block("4. Scenarios by number of tool calls, by industry", ["Industry", *TOOL_LEVELS.keys(), "Tool calls in total"],
          [[ind["industry"], *[sum(1 for s in ind["scenarios"] if tool_level(s) == t) for t in TOOL_LEVELS],
            sum(len(s.get("tools") or []) for s in ind["scenarios"])] for ind in inds]
          + [["All", *[sum(1 for s in all_s if tool_level(s) == t) for t in TOOL_LEVELS],
              sum(len(s.get("tools") or []) for s in all_s)]])
    block("5. Difficulty and direction by industry",
          ["Industry", "Easy", "Medium", "Hard", "Inbound", "Outbound"],
          [[ind["industry"], *[sum(1 for s in ind["scenarios"] if s["diff"] == d) for d in ("Easy", "Medium", "Hard")],
            sum(1 for s in ind["scenarios"] if s["dir"] == "in"), sum(1 for s in ind["scenarios"] if s["dir"] == "out")]
           for ind in inds])
    kinds = Counter()
    for s in all_s:
        for ln in s["_lines"]:
            kinds[ln["kind"]] += 1
            if ln["wait"]:
                kinds["silence"] += 1
            if ln["speaker"] != "customer":
                kinds["second"] += 1
    block("6. Planted moments across all scenarios", ["Moment", "Count"],
          [["Customer lines in total", sum(len(s["_lines"]) for s in all_s)],
           ["Interruptions (agent must stop)", kinds["cut"]],
           ["Acknowledgements (agent must keep talking)", kinds["bc"]],
           ["Mid-sentence pauses (agent must wait)", kinds["pause"]],
           ["Customer silences (agent checks in once)", kinds["silence"]],
           ["Lines by a second customer (phone handed over)", kinds["second"]],
           ["Scenarios with background sound", sum(1 for s in all_s if s.get("noise"))],
           ["… of which with a decoy voice", sum(1 for s in all_s if s.get("noise", {}) and s["noise"].get("decoy"))]])
    from statistics import median
    tl = [len(s["tools"]) for s in all_s]
    ll = [len(s["_lines"]) for s in all_s]
    block("7. Tool-call load", ["Measure", "Value"],
          [["Tool calls in total", sum(tl)],
           ["Tool calls per call — fewest / median / most", f"{min(tl)} / {median(tl):g} / {max(tl)}"],
           ["Customer lines per call — fewest / median / most", f"{min(ll)} / {median(ll):g} / {max(ll)}"],
           ["Tasks per call — median", f"{median(len(set(s['_task_of'])) for s in all_s):g}"],
           ["Calls with a tool error to recover from", sum(1 for s in all_s if any(is_error(c) for c in s["tools"]))],
           ["Tool calls that return an error", sum(is_error(c) for s in all_s for c in s["tools"])],
           ["Calls the agent must not make (one or more in every call)", sum(len(s.get("forbid") or []) for s in all_s)],
           ["Calls with an order rule", sum(1 for s in all_s if (s.get("checks") or {}).get("tool_order"))]])


def main() -> int:
    inds = load_sources()
    errs, seen = [], set()
    for ind in inds:
        for s in ind["scenarios"]:
            try:
                s["_lines"] = [parse_line(s.get("id", "?"), i, raw) for i, raw in enumerate(s["lines"], 1)]
            except ValueError as e:
                errs.append(str(e))
                s["_lines"] = []
                continue
            errs += validate(ind, s, seen)
            errs += validate_facts(s)
    plan_errs, plan_notes = check_plan(inds)
    errs += plan_errs
    if errs:
        print("\n".join(errs))
        return 1
    sys.path.insert(0, str(HERE))
    wb = build(inds)
    wb.save(OUT)
    all_s = [s for ind in inds for s in ind["scenarios"]]
    lc = Counter(s["lang"] for s in all_s)
    print(f"{len(all_s)} scenarios → {OUT.relative_to(HERE.parent.parent.parent)}")
    print("  " + " · ".join(f"{LANGS[k][0]} {lc[k]}" for k in LANGS))
    print("  " + " · ".join(f"{k} {sum(1 for s in all_s if tool_level(s) == k)}" for k in TOOL_LEVELS))
    print("  " + " · ".join(f"{a} {sum(1 for s in all_s if a in aspects(s))}" for a in ASPECTS))
    for n in plan_notes:
        print("  " + n)
    for ind in inds:
        c = Counter(s["lang"] for s in ind["scenarios"])
        print(f"  {ind['code']:5} {len(ind['scenarios']):3}/{PLANNED[ind['code']]}  " + " ".join(f"{k}={c[k]}" for k in LANGS)
              + f"  tools={sum(len(s['tools']) for s in ind['scenarios'])}"
              + f"  errors={sum(1 for s in ind['scenarios'] if any(is_error(x) for x in s['tools']))}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
