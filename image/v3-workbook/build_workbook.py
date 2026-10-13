"""Build the image v3 scenario workbook (.xlsx) from src/*.yaml.

    voice/.venv/bin/python image/v3-workbook/build_workbook.py
    voice/.venv/bin/python image/v3-workbook/build_workbook.py --check      # validate only
    voice/.venv/bin/python image/v3-workbook/build_workbook.py --src DIR --out FILE

Same pattern as voice/live/tool-heavy-workbook: YAML sources, a validating build,
and a workbook whose every count is derived from the sources.

    src/rubric.yaml         the rubric: criteria, what full and half marks look like, pass score
    src/sources.yaml        every input image, with a model-neutral source prompt
    src/<fam>.yaml          one header per family: rubric weights and gate, targets, slot grid
    src/<fam>-<ind>.yaml    part files: `code` and `scenarios`, merged in file order

The bank is written one industry at a time. Rules that can only hold for a full
family or a full bank are reported as "pending" until then; the build stops on
any rule that is already broken.
"""
from __future__ import annotations

import argparse
import re
import sys
from collections import Counter, OrderedDict, defaultdict
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
SRC = HERE / "src"
OUT = ROOT / "image" / "docs" / "sheets" / "2026-10-08-image-v3-scenarios.xlsx"
LANE = "image"
VIDEO_SRC = ROOT / "video" / "v3-workbook" / "src"

# ---------------------------------------------------------------- the plan

INDUSTRIES = OrderedDict([
    ("FLM", "Film, OTT & micro-drama"),
    ("ADS", "Advertising & agencies"),
    ("ECM", "E-commerce, D2C & retail"),
    ("GAM", "Gaming"),
    ("RES", "Real estate, interiors & architecture"),
    ("MED", "Media, publishing & events"),
])

FAMILY_ORDER = ["COV", "STB", "SHT", "CHR", "CMP", "TYP", "PRD", "EDT", "PHY"]
TASKS = OrderedDict([("t2i", "Text to image"), ("edit", "Image edit")])

# Slots per family and industry: (scenarios, t2i, edit). Industries not listed get none.
PLAN = {
    "COV": {"FLM": (15, 3, 12), "ADS": (6, 1, 5), "ECM": (6, 1, 5), "GAM": (6, 2, 4), "RES": (10, 3, 7), "MED": (2, 0, 2)},
    "STB": {"FLM": (14, 4, 10), "ADS": (8, 3, 5), "ECM": (2, 1, 1), "GAM": (6, 2, 4)},
    "SHT": {"FLM": (8, 3, 5), "ADS": (7, 3, 4), "ECM": (5, 2, 3), "GAM": (5, 2, 3)},
    "CHR": {"FLM": (16, 4, 12), "ADS": (4, 1, 3), "ECM": (5, 1, 4), "GAM": (10, 4, 6)},
    "CMP": {"FLM": (10, 9, 1), "ADS": (7, 6, 1), "ECM": (6, 5, 1), "GAM": (3, 3, 0), "RES": (3, 2, 1), "MED": (6, 5, 1)},
    "TYP": {"FLM": (4, 3, 1), "ADS": (8, 6, 2), "ECM": (8, 6, 2), "GAM": (4, 3, 1), "MED": (16, 12, 4)},
    "PRD": {"ADS": (10, 2, 8), "ECM": (20, 2, 18), "GAM": (3, 1, 2), "RES": (2, 0, 2)},
    "EDT": {"FLM": (3, 0, 3), "ADS": (4, 0, 4), "ECM": (8, 0, 8), "GAM": (4, 0, 4), "RES": (4, 0, 4), "MED": (2, 0, 2)},
    "PHY": {"FLM": (5, 3, 2), "ADS": (6, 4, 2), "GAM": (4, 3, 1), "RES": (11, 7, 4), "MED": (4, 3, 1)},
}
# Extreme scenarios per family (about 40%); the rest are hard.
EXTREME = {"COV": 18, "STB": 12, "SHT": 10, "CHR": 14, "CMP": 14, "TYP": 16, "PRD": 14, "EDT": 10, "PHY": 12}
BANK_TOTAL = 300
MIN_SHARE = {"t2i": 0.35, "edit": 0.50}
MIN_INDIA = 0.30
MAX_SUBJECT_SHARE = 0.25
MIN_PER_FAMILY = {"industries": 4, "subjects": 6, "styles": 3, "lights": 4}
MAX_KIT_USES = 5
CHECKS_RANGE = (12, 25)
MEMBER_CHECKS_RANGE = (2, 5)
MIN_LEVERS = 3
MIN_LEVER_FAMILIES = 3
# Levers that cannot reach MIN_LEVER_FAMILIES by design: only the set families produce sets.
LEVER_FAMILY_EXEMPT = {"set": "only Multi-angle coverage and Character continuity produce sets"}
MIN_PROMPT_WORDS = 80
MIN_SOURCE_PROMPT_WORDS = 60

SUBJECTS = OrderedDict([
    ("person", "Single person"), ("group", "Group (3+)"), ("child-elderly", "Child or elderly"),
    ("product", "Product"), ("food", "Food"), ("fashion", "Fashion"), ("vehicle", "Vehicle"),
    ("animal", "Animal"), ("interior", "Interior"), ("architecture", "Architecture"),
    ("landscape", "Landscape"), ("stylised", "Stylised character"),
])
STYLES = OrderedDict([
    ("photoreal", "Photoreal"), ("cinematic", "Cinematic film"), ("illustration", "Illustration or anime"),
    ("3d", "3D render"), ("sketch", "Pencil or sketch"), ("graphic", "Graphic design"),
])
LIGHTS = OrderedDict([
    ("daylight", "Daylight"), ("golden", "Golden hour"), ("night", "Night"), ("neon", "Neon or mixed"),
    ("practical", "Indoor practical"), ("fog-rain", "Fog or rain"), ("noon", "Harsh noon"),
])
LOCALES = OrderedDict([("IN", "Indian context"), ("global", "Global")])
SCRIPTS = OrderedDict([
    ("none", "No text"), ("latin", "Latin"), ("devanagari", "Devanagari"),
    ("telugu", "Telugu"), ("tamil", "Tamil"), ("mixed", "Mixed scripts"),
])
TEXT_SCRIPTS = ("latin", "devanagari", "telugu", "tamil")
SCRIPT_RANGES = {"devanagari": "ऀ-ॿ", "telugu": "ఀ-౿", "tamil": "஀-௿"}
DIFFS = OrderedDict([("hard", "Hard"), ("extreme", "Extreme")])

# Hardness levers: what makes a scenario hard. Each scenario uses 3 or more.
LEVERS = OrderedDict([
    ("set", "Several outputs that must agree with each other"),
    ("unseen", "Has to show what the input never shows (reverse side, back, behind the camera)"),
    ("count", "Exact numbers: people, objects, panels, beats"),
    ("layout", "Exact positions: who is left or right, what is where"),
    ("exact-text", "Exact words that must appear letter for letter"),
    ("script", "Text in Devanagari, Telugu or Tamil"),
    ("small-text", "Small or dense text that must stay readable"),
    ("exclude", "Named things that must not appear"),
    ("off-centre", "Composition against the default centred framing"),
    ("camera", "A named lens, shot size, camera height or angle"),
    ("light-lock", "A light direction or time of day that must hold"),
    ("physics", "Reflections, shadows, refraction, occlusion or contact"),
    ("identity", "A face, character or product kept from a reference"),
    ("multi-ref", "Two to four inputs combined"),
    ("locale", "Indian cultural detail that a default model gets wrong"),
    ("keep-rest", "Everything outside the change stays the same"),
    ("multi-step", "Several changes in one instruction"),
    ("long-brief", "A long brief with ten or more constraints"),
    ("continuity", "Screen direction, props or costume carried across frames"),
    ("time-code", "Time codes or durations that must add up"),
])

# Rubric criteria that only make sense with an input image.
INPUT_CRITERIA = ("reference_fidelity", "preservation")

SET_MODES = OrderedDict([
    ("separate", "Separate calls, each from the same inputs"),
    ("chained", "Each call also gets the earlier outputs"),
    ("one-sheet", "One image holding every member as a panel"),
])
RUNNABLE = OrderedDict([("yes", "Yes"), ("partly", "Partly"), ("no", "No")])
INPUT_ROLES = ("source", "reference")

# Sahil's six example tests (screen-share of 6 October 2026).
SAHIL = OrderedDict([
    ("composition", ("Composition and texture against a composition-led tool", "223–231", "CMP")),
    ("product-bible", ("Production bible built from one product photo", "232–233", "SHT")),
    ("storyboard", ("Storyboard built from a base image", "234–235", "STB")),
    ("beat-sheet", ("Time-coded beat sheet", "236–237", "SHT")),
    ("typography", ("Typography posters against a type-led tool", "238–239", "TYP")),
    ("angles", ("Many camera angles from one shot", "239–244", "COV")),
])

# Source and scenario prompts must read the same whichever tool runs them.
VENDOR_RE = re.compile(
    r"\b(gemini|nano[\s-]?banana|nb[\s-]?pro|imagen|veo|gpt|chatgpt|dall[\s·\-]?e|sora|midjourney|"
    r"recraft|flux|seedance|kling|runway[\s-]?(ml|gen)|stable[\s-]?diffusion|sdxl|firefly|ideogram|"
    r"leonardo[\s.]?ai|pika|hailuo|luma|openai|google|deepmind|bytedance|adobe|elevenlabs)\b",
    re.I)
TOOL_SYNTAX_RE = re.compile(
    r"(--[a-z]{1,12}\b|::\s*-?\d|\(\s*[\w\s-]+:\s*\d+(\.\d+)?\s*\)|\[\[|\]\]|\{\{|\}\})", re.I)

SLOT_ID_RE = re.compile(r"^(?P<fam>[A-Z]{3})-(?P<ind>[A-Z]{3})-(?P<n>\d{2})$")
SOURCE_ID_RE = re.compile(r"^SRC-[A-Z0-9]+(-[A-Z0-9]+)*$")


# ---------------------------------------------------------------- loading

def load(src: Path):
    """Return (families by code, sources by id)."""
    fams, parts = {}, []
    for f in sorted(src.glob("*.yaml")):
        if f.name.startswith("_") or f.name in ("sources.yaml", "rubric.yaml", "customers.yaml"):
            continue
        data = yaml.safe_load(f.read_text(encoding="utf-8")) or {}
        data["_file"] = f.name
        if "family" in data:
            data["slots"] = list(data.get("slots") or [])
            data["scenarios"] = list(data.get("scenarios") or [])
            for s in data["scenarios"]:
                s["_file"] = f.name
            fams[data["code"]] = data
        else:
            parts.append(data)
    for p in parts:
        if p.get("code") not in fams:
            raise SystemExit(f"{p['_file']}: no family header for code {p.get('code')}")
        for s in p.get("scenarios") or []:
            s["_file"] = p["_file"]
            fams[p["code"]]["scenarios"].append(s)
    missing = [c for c in FAMILY_ORDER if c not in fams]
    if missing:
        raise SystemExit(f"missing family headers: {', '.join(missing)}")
    srcs_file = src / "sources.yaml"
    sources = OrderedDict()
    for s in (yaml.safe_load(srcs_file.read_text(encoding="utf-8")) or {}).get("sources") or []:
        if s.get("id") in sources:
            raise SystemExit(f"sources.yaml: duplicate source {s.get('id')}")
        sources[s["id"]] = s
    return OrderedDict((c, fams[c]) for c in FAMILY_ORDER), sources


def load_customers(path: Path, types_path: Path = None) -> tuple[dict, dict]:
    """(types by code, customer codes by scenario id) from customers.yaml. The
    types live in the image bank's file; the video bank reads them from there."""
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    types = (yaml.safe_load(types_path.read_text(encoding="utf-8")) if types_path else data).get("types") or {}
    for k, v in types.items():
        if not (v or {}).get("name") or not v.get("detail"):
            raise SystemExit(f"customers.yaml: type {k} needs a name and a detail")
    return types, data.get("scenarios") or {}


def validate_customers(r, written: dict, customers: tuple[dict, dict]) -> None:
    """Every scenario names 1–3 known customer types; no stray ids."""
    types, by_id = customers
    for sid in written:
        got = by_id.get(sid) or []
        if not 1 <= len(got) <= 3:
            r.err(f"{sid}: customers.yaml names {len(got)} target customers; a scenario has 1–3")
        for c in got:
            if c not in types:
                r.err(f"{sid}: unknown customer type '{c}' in customers.yaml")
        if len(set(got)) != len(got):
            r.err(f"{sid}: a target customer is named twice in customers.yaml")
    for sid in by_id:
        if sid not in written:
            r.err(f"customers.yaml: {sid} is not a scenario in the bank")


def customers_text(customers: tuple[dict, dict], sid: str) -> str:
    types, by_id = customers
    return ", ".join(types[c]["name"] for c in by_id[sid])


def customers_sheet(wb, scen: list, customers: tuple[dict, dict], table) -> None:
    types, by_id = customers
    ws = wb.create_sheet("Target customers")
    n = Counter(c for s in scen for c in by_id.get(s["id"], []))
    rows = [[v["name"], v["detail"], n[k]] for k, v in types.items() if n[k]]
    table(ws, ["Target customer", "Who it covers", "Scenarios"], rows, [32, 80, 10])


def load_rubric(path: Path, lane: str = LANE) -> dict:
    """The rubric: scoring rules, pass score and this lane's criteria by key."""
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    crit = OrderedDict()
    for c in data.get("criteria") or []:
        for k in ("key", "name", "lanes", "judges", "full", "half", "applies"):
            if not c.get(k):
                raise SystemExit(f"{path.name}: criterion {c.get('key', '?')} is missing {k}")
        if lane in c["lanes"]:
            crit[c["key"]] = c
    return {"scoring": data["scoring"], "pass_score": float(data["pass_score"]), "criteria": crit}


def words(text: str) -> int:
    return len(re.findall(r"\w+", text or ""))


def banned(text: str) -> list[str]:
    hits = [m.group(0) for m in VENDOR_RE.finditer(text or "")]
    hits += [m.group(0) for m in TOOL_SYNTAX_RE.finditer(text or "")]
    return hits


def has_script(text: str, script: str) -> bool:
    if script == "latin":
        return bool(re.search(r"[A-Za-z]", text))
    return bool(re.search(f"[{SCRIPT_RANGES[script]}]", text))


# ---------------------------------------------------------------- validation

class Report:
    def __init__(self):
        self.errors: list[str] = []
        self.rules: list[tuple[str, str, str, str]] = []   # (scope, rule, value, status)

    def err(self, msg: str):
        self.errors.append(msg)

    def rule(self, scope: str, rule: str, value, status: str):
        self.rules.append((scope, rule, str(value), status))


def validate(fams: OrderedDict, sources: OrderedDict, rubric: dict, customers: tuple) -> Report:
    r = Report()
    slots, seen = {}, set()

    # --- family rubrics
    for code, fam in fams.items():
        fr = fam.get("rubric") or {}
        w, gate = fr.get("weights") or {}, fr.get("gate") or []
        if not w or not gate or not fr.get("gate_rule"):
            r.err(f"{code}: the family rubric needs weights, a gate and a gate_rule")
        for k, v in w.items():
            if k not in rubric["criteria"]:
                r.err(f"{code}: rubric weight for unknown criterion '{k}'")
            if not (isinstance(v, int) and v > 0):
                r.err(f"{code}: rubric weight for {k} must be a whole number above 0")
        for k in gate:
            if k not in w:
                r.err(f"{code}: gate criterion '{k}' has no weight")

    # --- slots
    for code, fam in fams.items():
        for t in fam.get("targets") or {}:
            if not re.match(r"^[a-z][a-z-]*$", t):
                r.err(f"{code}: target '{t}' must be lower-case words joined by hyphens")
        for sl in fam["slots"]:
            sid = sl.get("id", "?")
            m = SLOT_ID_RE.match(sid)
            if not m:
                r.err(f"{sid}: slot id must look like COV-FLM-01")
                continue
            if sid in seen:
                r.err(f"{sid}: duplicate slot id")
            seen.add(sid)
            if m["fam"] != code:
                r.err(f"{sid}: slot is in {fam['_file']} but its id says family {m['fam']}")
            if m["ind"] not in INDUSTRIES:
                r.err(f"{sid}: unknown industry {m['ind']}")
            elif m["ind"] not in PLAN[code]:
                r.err(f"{sid}: the plan has no {INDUSTRIES[m['ind']]} slots in {code}")
            for key, allowed in (("task", TASKS), ("subject", SUBJECTS), ("style", STYLES),
                                 ("light", LIGHTS), ("locale", LOCALES), ("diff", DIFFS)):
                if sl.get(key) not in allowed:
                    r.err(f"{sid}: {key} '{sl.get(key)}' is not one of {', '.join(allowed)}")
            sl.setdefault("script", "none")
            if sl["script"] not in SCRIPTS:
                r.err(f"{sid}: script '{sl['script']}' is not one of {', '.join(SCRIPTS)}")
            if fam.get("text_family") and sl["script"] == "none":
                r.err(f"{sid}: {code} is a text family, so every slot names its script")
            if sl.get("target") not in (fam.get("targets") or {}):
                r.err(f"{sid}: target '{sl.get('target')}' is not a {code} target")
            if not sl.get("idea"):
                r.err(f"{sid}: missing idea")
            srcs = sl.get("sources") or []
            sl["sources"] = srcs
            if sl.get("task") == "edit" and not srcs:
                r.err(f"{sid}: an edit slot names at least one source")
            if sl.get("task") == "t2i" and srcs:
                r.err(f"{sid}: a text-to-image slot has no sources")
            if len(srcs) > 4:
                r.err(f"{sid}: at most four inputs")
            for s in srcs:
                if s not in sources:
                    r.err(f"{sid}: source {s} is not in sources.yaml")
            sl["_fam"], sl["_ind"] = code, m["ind"]
            slots[sid] = sl

    # --- slot plan per family and industry (exact once an industry is started)
    for code, fam in fams.items():
        by_ind = defaultdict(list)
        for sl in fam["slots"]:
            by_ind[sl.get("_ind")].append(sl)
        for ind, (qty, t2i, edit) in PLAN[code].items():
            got = by_ind.get(ind, [])
            if not got:
                continue
            tasks = Counter(s["task"] for s in got)
            if len(got) != qty:
                r.err(f"{code} × {INDUSTRIES[ind]}: {len(got)} slots, the plan says {qty}")
            if (tasks["t2i"], tasks["edit"]) != (t2i, edit):
                r.err(f"{code} × {INDUSTRIES[ind]}: task split {tasks['t2i']} t2i / {tasks['edit']} edit, "
                      f"the plan says {t2i} / {edit}")
            nums = sorted(int(s["id"][-2:]) for s in got)
            if nums != list(range(1, len(got) + 1)):
                r.err(f"{code} × {INDUSTRIES[ind]}: slot numbers must run 01–{len(got):02d} with no gaps")

    # --- near-duplicates: family + subject + industry + failure target
    keys = defaultdict(list)
    for sl in slots.values():
        keys[(sl["_fam"], sl["subject"], sl["_ind"], sl["target"])].append(sl["id"])
    for k, ids in keys.items():
        if len(ids) > 1:
            r.err(f"{', '.join(ids)}: same family, subject, industry and failure target ({k[1]}, {k[3]})")

    # --- running caps against the planned family size
    for code, fam in fams.items():
        planned = sum(q for q, _, _ in PLAN[code].values())
        subj = Counter(sl["subject"] for sl in fam["slots"])
        cap = int(planned * MAX_SUBJECT_SHARE)
        for s, n in subj.items():
            if n > cap:
                r.err(f"{code}: {n} '{s}' slots; no subject may pass {cap} ({MAX_SUBJECT_SHARE:.0%} of {planned})")
        diffs = Counter(sl["diff"] for sl in fam["slots"])
        if diffs["extreme"] > EXTREME[code]:
            r.err(f"{code}: {diffs['extreme']} extreme slots, the plan allows {EXTREME[code]}")
        if diffs["hard"] > planned - EXTREME[code]:
            r.err(f"{code}: {diffs['hard']} hard slots, the plan allows {planned - EXTREME[code]}")

    # --- sources
    uses = defaultdict(list)
    for sl in slots.values():
        for s in sl["sources"]:
            uses[s].append(sl["id"])
    for sid, s in sources.items():
        if not SOURCE_ID_RE.match(sid):
            r.err(f"{sid}: source ids look like SRC-KIDS-TREE")
        if s.get("kind") != "image":
            r.err(f"{sid}: kind must be image in the image lane")
        if s.get("role") not in ("scene", "reference"):
            r.err(f"{sid}: role must be scene or reference")
        for k in ("source_prompt", "aspect", "resolution"):
            if not s.get(k):
                r.err(f"{sid}: missing {k}")
        if not re.match(r"^\d+(\.\d+)?:\d+$", str(s.get("aspect", ""))):
            r.err(f"{sid}: aspect looks like 16:9")
        if words(s.get("source_prompt", "")) < MIN_SOURCE_PROMPT_WORDS:
            r.err(f"{sid}: source prompt under {MIN_SOURCE_PROMPT_WORDS} words; describe layout, light and lens")
        if len(s.get("must_contain") or []) < 2:
            r.err(f"{sid}: must_contain needs at least two acceptance items")
        if not s.get("avoid"):
            r.err(f"{sid}: avoid needs at least one item")
        text = " ".join([s.get("source_prompt", "")] + list(s.get("avoid") or []) + list(s.get("must_contain") or []))
        hits = banned(text)
        if hits:
            r.err(f"{sid}: source prompt names a model or vendor, or uses tool syntax: {', '.join(sorted(set(hits)))}")
        declared = sorted(s.get("used_by") or [])
        if declared != sorted(uses.get(sid, [])):
            r.err(f"{sid}: used_by {declared} does not match the slots that name it {sorted(uses.get(sid, []))}")
        if len(uses.get(sid, [])) > MAX_KIT_USES:
            r.err(f"{sid}: used by {len(uses[sid])} slots; the cap is {MAX_KIT_USES}")

    # --- scenarios
    written = {}
    for code, fam in fams.items():
        for s in fam["scenarios"]:
            sid = s.get("id", "?")
            where = f"{sid} ({s.get('_file')})"
            if sid not in slots:
                r.err(f"{where}: no slot with this id")
                continue
            sl = slots[sid]
            if sl["_fam"] != code:
                r.err(f"{where}: filed under {code} but the slot is {sl['_fam']}")
            if sid in written:
                r.err(f"{where}: written twice")
            written[sid] = s
            s["_slot"] = sl
            validate_scenario(r, where, s, sl, fam, rubric)

    # --- coverage rules (pending until the family or bank is full)
    coverage(r, fams, slots, written, sources, uses)
    validate_customers(r, written, customers)
    return r


def validate_scenario(r: Report, where: str, s: dict, sl: dict, fam: dict, rubric: dict) -> None:
    for k in ("title", "persona", "evidence", "levers", "prompt", "deliverable", "checks",
              "competitor", "runnable_today"):
        if not s.get(k):
            r.err(f"{where}: missing {k}")
    levers = s.get("levers") or []
    for lv in levers:
        if lv not in LEVERS:
            r.err(f"{where}: unknown lever '{lv}'")
    if len(set(levers)) < MIN_LEVERS:
        r.err(f"{where}: {len(set(levers))} levers; each scenario uses at least {MIN_LEVERS}")
    if words(s.get("prompt", "")) < MIN_PROMPT_WORDS:
        r.err(f"{where}: prompt under {MIN_PROMPT_WORDS} words")

    # inputs must be exactly the slot's sources
    inputs = s.get("inputs") or []
    if sl["task"] == "edit":
        got = sorted(i.get("source") for i in inputs)
        if got != sorted(sl["sources"]):
            r.err(f"{where}: inputs {got} do not match the slot's sources {sorted(sl['sources'])}")
        roles = Counter(i.get("role") for i in inputs)
        if any(x not in INPUT_ROLES for x in roles):
            r.err(f"{where}: input roles are source or reference")
        if roles["source"] != 1:
            r.err(f"{where}: an edit has exactly one input with role source")
    elif inputs:
        r.err(f"{where}: a text-to-image scenario has no inputs")

    # checks
    checks = s.get("checks") or []
    lo, hi = CHECKS_RANGE
    if not lo <= len(checks) <= hi:
        r.err(f"{where}: {len(checks)} checks; a scenario has {lo}–{hi}")
    crits = Counter()
    for i, c in enumerate(checks, 1):
        if not (isinstance(c, list) and len(c) == 2):
            r.err(f"{where}: check {i} must be [criterion, question]")
            continue
        if c[0] not in rubric["criteria"]:
            r.err(f"{where}: check {i} has unknown rubric criterion '{c[0]}'")
        if not str(c[1]).rstrip().endswith("?"):
            r.err(f"{where}: check {i} must be a yes/no question ending in '?'")
        crits[c[0]] += 1

    # sets
    sm = s.get("set_members")
    rng = fam.get("set")
    all_prompts = [s.get("prompt", "")]
    if rng:
        if not sm:
            r.err(f"{where}: {sl['_fam']} scenarios are sets; add set_members")
        else:
            if sm.get("mode") not in SET_MODES:
                r.err(f"{where}: set mode must be one of {', '.join(SET_MODES)}")
            if len(sm.get("must_stay_same") or []) < 3:
                r.err(f"{where}: must_stay_same needs at least three items")
            members = sm.get("members") or []
            if not rng["min"] <= len(members) <= rng["max"]:
                r.err(f"{where}: {len(members)} set members; {sl['_fam']} sets have {rng['min']}–{rng['max']}")
            mlo, mhi = MEMBER_CHECKS_RANGE
            for j, mem in enumerate(members, 1):
                if not mem.get("name") or not mem.get("prompt"):
                    r.err(f"{where}: member {j} needs a name and a prompt")
                all_prompts.append(mem.get("prompt", ""))
                mc = mem.get("checks") or []
                if not mlo <= len(mc) <= mhi:
                    r.err(f"{where}: member {j} has {len(mc)} checks; members have {mlo}–{mhi}")
                for k, c in enumerate(mc, 1):
                    if not (isinstance(c, list) and len(c) == 2 and c[0] in rubric["criteria"]
                            and str(c[1]).rstrip().endswith("?")):
                        r.err(f"{where}: member {j} check {k} must be [criterion, yes/no question?]")
                    else:
                        crits[c[0]] += 1
            if not crits["set_consistency"]:
                r.err(f"{where}: a set needs at least one set_consistency check")
            if "set" not in levers:
                r.err(f"{where}: a set scenario uses the 'set' lever")
    elif sm:
        r.err(f"{where}: {sl['_fam']} scenarios are single outputs; remove set_members")

    # model-neutral prompts
    hits = banned(" ".join(all_prompts))
    if hits:
        r.err(f"{where}: prompt names a model or vendor, or uses tool syntax: {', '.join(sorted(set(hits)))}")

    # exact text: given in the prompt, checked, and in the slot's script
    exact = s.get("exact_text") or []
    joined = " ".join(" ".join(all_prompts).split())   # line wraps in the YAML do not count
    for t in exact:
        if " ".join(str(t).split()) not in joined:
            r.err(f"{where}: exact text '{t}' is not given word for word in the prompt")
    if exact and not crits["text_accuracy"]:
        r.err(f"{where}: exact text needs at least one text_accuracy check")

    # rubric: every criterion checked is weighted by the family, and the gate applies
    fr = fam.get("rubric") or {}
    if not crits["prompt_adherence"]:
        r.err(f"{where}: every scenario has at least one prompt_adherence check")
    for k in crits:
        if k in rubric["criteria"] and k not in (fr.get("weights") or {}):
            r.err(f"{where}: checks {k}, which the {sl['_fam']} rubric does not weight")
        if k in INPUT_CRITERIA and not inputs:
            r.err(f"{where}: {k} checks need an input image")
    if fr.get("gate") and not any(crits[k] for k in fr["gate"]):
        r.err(f"{where}: checks none of the gate criteria ({', '.join(fr['gate'])})")
    s["_crits"] = crits
    script = sl["script"]
    if script in SCRIPT_RANGES and not any(has_script(t, script) for t in exact):
        r.err(f"{where}: the slot's script is {script} but no exact text uses it")
    if script == "mixed" and sum(any(has_script(t, sc) for t in exact) for sc in TEXT_SCRIPTS) < 2:
        r.err(f"{where}: script is mixed but the exact text uses fewer than two scripts")
    if script == "none" and exact:
        r.err(f"{where}: the slot has no text but the scenario lists exact text; set the slot's script")

    rt = s.get("runnable_today") or {}
    if isinstance(rt.get("status"), bool):
        r.err(f"{where}: write runnable_today status in quotes (\"yes\"); YAML reads a bare yes as true")
    elif rt.get("status") not in RUNNABLE or not rt.get("note"):
        r.err(f"{where}: runnable_today needs status (yes, partly, no) and a note")
    if s.get("sahil") and s["sahil"] not in SAHIL:
        r.err(f"{where}: sahil must be one of {', '.join(SAHIL)}")
    if s.get("sahil") and not s.get("seed"):
        r.err(f"{where}: a recreation of Sahil's example is a seed")


def coverage(r: Report, fams, slots, written, sources, uses) -> None:
    all_slots = list(slots.values())
    bank_full = len(all_slots) == BANK_TOTAL
    for code, fam in fams.items():
        planned = sum(q for q, _, _ in PLAN[code].values())
        fs = fam["slots"]
        full = len(fs) == planned
        scope = f"{code} {fam['family']}"
        st = lambda ok: ("met" if ok else "broken") if full else "pending"

        def chk(rule, value, ok):
            r.rule(scope, rule, value, st(ok))
            if full and not ok:
                r.err(f"{code}: {rule} — {value}")

        chk("Slots declared", f"{len(fs)} of {planned}", len(fs) == planned)
        inds = {s["_ind"] for s in fs}
        chk(f"At least {MIN_PER_FAMILY['industries']} industries", len(inds), len(inds) >= MIN_PER_FAMILY["industries"])
        subj = Counter(s["subject"] for s in fs)
        chk(f"At least {MIN_PER_FAMILY['subjects']} subjects", len(subj), len(subj) >= MIN_PER_FAMILY["subjects"])
        sty = {s["style"] for s in fs}
        chk(f"At least {MIN_PER_FAMILY['styles']} styles", len(sty), len(sty) >= MIN_PER_FAMILY["styles"])
        lig = {s["light"] for s in fs}
        chk(f"At least {MIN_PER_FAMILY['lights']} light conditions", len(lig), len(lig) >= MIN_PER_FAMILY["lights"])
        top = subj.most_common(1)[0] if subj else ("—", 0)
        share = top[1] / max(len(fs), 1)
        chk(f"No subject above {MAX_SUBJECT_SHARE:.0%}", f"{top[0]} {share:.0%}", share <= MAX_SUBJECT_SHARE)
        ext = sum(1 for s in fs if s["diff"] == "extreme")
        chk("Extreme share about 40%", f"{ext} of {len(fs)} (plan {EXTREME[code]})", ext == EXTREME[code])
        t = Counter(s["task"] for s in fs)
        pt = sum(v[1] for v in PLAN[code].values()), sum(v[2] for v in PLAN[code].values())
        chk("Task split matches the plan", f"{t['t2i']} t2i / {t['edit']} edit (plan {pt[0]} / {pt[1]})",
            (t["t2i"], t["edit"]) == pt)
        if fam.get("text_family"):
            scr = {s["script"] for s in fs}
            missing = [x for x in TEXT_SCRIPTS if x not in scr]
            chk("All four scripts", ", ".join(sorted(scr)) or "—", not missing)
        n_written = sum(1 for s in fs if s["id"] in written)
        r.rule(scope, "Scenarios written", f"{n_written} of {planned}", "met" if n_written == planned else "pending")
        seeds = [s for s in fam["scenarios"] if s.get("seed")]
        r.rule(scope, "Seeds (3 per family)", len(seeds), "met" if len(seeds) == 3 else ("broken" if len(seeds) > 3 else "pending"))
        if len(seeds) > 3:
            r.err(f"{code}: {len(seeds)} seeds; a family has 3")

    scope = "Whole bank"
    n = len(all_slots)
    india = sum(1 for s in all_slots if s["locale"] == "IN")
    ok = india / max(n, 1) >= MIN_INDIA
    r.rule(scope, f"At least {MIN_INDIA:.0%} Indian context", f"{india} of {n} ({india / max(n, 1):.0%})",
           ("met" if ok else "broken") if bank_full else ("on track" if ok else "behind"))
    if bank_full and not ok:
        r.err(f"bank: Indian context {india / n:.0%}, needs {MIN_INDIA:.0%}")
    tasks = Counter(s["task"] for s in all_slots)
    for task, share in MIN_SHARE.items():
        got = tasks[task] / max(n, 1)
        ok = got >= share
        r.rule(scope, f"{TASKS[task]} at least {share:.0%}", f"{tasks[task]} of {n} ({got:.0%})",
               ("met" if ok else "broken") if bank_full else ("on track" if ok else "behind"))
        if bank_full and not ok:
            r.err(f"bank: {TASKS[task]} share {got:.0%}, needs {share:.0%}")
    all_written = bank_full and len(written) == BANK_TOTAL
    lever_fams = defaultdict(set)
    for s in written.values():
        for lv in s.get("levers") or []:
            lever_fams[lv].add(s["_slot"]["_fam"])
    thin = [lv for lv in LEVERS if lv not in LEVER_FAMILY_EXEMPT and len(lever_fams[lv]) < MIN_LEVER_FAMILIES]
    r.rule(scope, f"Every lever in at least {MIN_LEVER_FAMILIES} families (except: "
           + "; ".join(f"{k}, {v}" for k, v in LEVER_FAMILY_EXEMPT.items()) + ")",
           "all" if not thin else f"{len(thin)} below: {', '.join(thin)}",
           ("met" if not thin else "broken") if all_written else "pending")
    if all_written and thin:
        r.err(f"bank: levers in fewer than {MIN_LEVER_FAMILIES} families: {', '.join(thin)}")
    unused = [sid for sid in sources if not uses.get(sid)]
    r.rule(scope, "Every source used by a slot", "all" if not unused else ", ".join(unused),
           "met" if not unused else "broken")
    for sid in unused:
        r.err(f"{sid}: no slot uses this source")
    sahil_done = {s.get("sahil") for s in written.values() if s.get("sahil")}
    r.rule(scope, "Sahil's six examples as seeds", f"{len(sahil_done)} of 6",
           "met" if len(sahil_done) == 6 else "pending")


# ---------------------------------------------------------------- derived values

def outputs_of(s: dict) -> int:
    sm = s.get("set_members")
    if not sm:
        return 1
    return 1 if sm.get("mode") == "one-sheet" else len(sm.get("members") or [])


def planned_outputs(code: str) -> str:
    return {"COV": "5–8", "CHR": "3–6"}.get(code, "1")


def rubric_of(s: dict, fam: dict) -> list:
    """[(criterion, weight share, checks)] for the criteria this scenario checks."""
    w = fam["rubric"]["weights"]
    crits = s["_crits"]
    total = sum(w[k] for k in crits)
    return [(k, w[k] / total, crits[k]) for k in w if k in crits]


def rubric_text(s: dict, fam: dict, rubric: dict) -> str:
    gate = fam["rubric"]["gate"]
    return "\n".join(f"{rubric['criteria'][k]['name']}: {share:.0%} · {n} check{'s' if n > 1 else ''}"
                     + (" · gate" if k in gate else "") for k, share, n in rubric_of(s, fam))


def pass_text(s: dict, fam: dict, rubric: dict) -> str:
    gate = [rubric["criteria"][k]["name"] for k in fam["rubric"]["gate"] if s["_crits"][k]]
    n = sum(s["_crits"][k] for k in fam["rubric"]["gate"])
    first = (f"The {gate[0]} check is a yes" if n == 1
             else f"All {n} {' and '.join(gate)} checks are yes")
    return f"{first}, and the score is {rubric['pass_score']:.1f} or more out of 10."


def target_label(fam: dict, t: str) -> str:
    return f"{t} — {(fam.get('targets') or {}).get(t, '')}"


def all_checks(s: dict) -> list:
    return list(s["checks"]) + [c for m in (s.get("set_members") or {}).get("members") or []
                                for c in m.get("checks") or []]


def members_text(s: dict) -> str:
    sm = s.get("set_members") or {}
    return "\n".join(f"{j}. {m['name']}: {m['prompt'].strip()}"
                     for j, m in enumerate(sm.get("members") or [], 1)) or "—"


def video_links() -> dict:
    """Image scenario id -> [(video scenario id, idea)], from the video bank's from_image."""
    links = defaultdict(list)
    for f in sorted(VIDEO_SRC.glob("v*.yaml")):
        for sl in (yaml.safe_load(f.read_text(encoding="utf-8")) or {}).get("slots") or []:
            if sl.get("from_image"):
                links[sl["from_image"]].append((sl["id"], sl["idea"]))
    return links


def video_text(links: dict, sid: str) -> str:
    return "\n".join(f"{v}: {idea}" for v, idea in links.get(sid, [])) or "—"


def inputs_text(s: dict) -> str:
    return "\n".join(f"{i['source']} ({i['role']})" + (f": {i['note']}" if i.get("note") else "")
                     for i in s.get("inputs") or []) or "—"


# ---------------------------------------------------------------- workbook helpers

HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
SECTION_FONT = Font(bold=True, size=11, color="1F3A5F")
WRAP = Alignment(wrap_text=True, vertical="top")
STATUS_FILL = {
    "met": PatternFill("solid", fgColor="D9EAD3"), "on track": PatternFill("solid", fgColor="D9EAD3"),
    "pending": PatternFill("solid", fgColor="FFF2CC"), "behind": PatternFill("solid", fgColor="FCE5CD"),
    "broken": PatternFill("solid", fgColor="F4CCCC"),
}


def table(ws, headers, rows, widths, freeze="B2"):
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, WRAP
    for row in rows:
        ws.append(row)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions


# ---------------------------------------------------------------- sheets

def build(fams: OrderedDict, sources: OrderedDict, rubric: dict, rep: Report, customers: tuple) -> Workbook:
    slots = [sl for f in fams.values() for sl in f["slots"]]
    scen = [s for f in fams.values() for s in sorted(f["scenarios"], key=lambda x: x["id"])]
    wb = Workbook()
    vlinks = video_links()

    # 1 · How to read
    ws = wb.active
    ws.title = "How to read"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 120
    ws.append(["Image scenario bank v3 — 300 scenarios, written to the 98th–99th percentile bar"])
    ws["A1"].font = TITLE_FONT
    outs = sum(outputs_of(s) for s in scen)
    started = [INDUSTRIES[i] for i in INDUSTRIES if any(sl["_ind"] == i for sl in slots)]
    paras = [
        ("What this is", "Hard image-generation and image-editing scenarios across 9 families and 6 industries. "
         "One scenario is one brief and one answer key, judged as a whole: a six-angle coverage set is one "
         "scenario with six outputs. Every input image has a model-neutral source prompt (Sources tab)."),
        ("Progress", f"{len(slots)} of {BANK_TOTAL} slots declared; {len(scen)} scenarios written "
         f"({sum(1 for s in scen if s.get('seed'))} seeds), {outs} outputs per model so far. "
         f"Industries started: {', '.join(started) or 'none'}. The bank is written one industry at a time."),
        ("Slots", "Each scenario fills a slot. The slot fixes family, industry, task, subject, style, light, "
         "locale, script, difficulty, failure target and inputs, so diversity is designed in rather than "
         "checked afterwards. Empty slots show what is still to write."),
        ("Sets", "A set is one scenario with several outputs. Each output is a separate request: the shared "
         "Prompt plus one line from 'What each output is'. A one-sheet set is a single request for one image "
         "holding every member as a panel. The set is judged as a whole, so the outputs must agree."),
        ("Target customers", "The kinds of business each scenario is for: where its persona works, then the "
         "brand category it serves. One to three per scenario, from a fixed list (Target customers tab)."),
        ("Review tab", "Persona, failure target and levers are for reviewing the bank, not for running it, "
         "so they sit on the Review tab."),
        ("Rubric", f"Every scenario is scored against one rubric of {len(rubric['criteria'])} criteria (Rubric tab). "
         f"Each scenario has {CHECKS_RANGE[0]}–{CHECKS_RANGE[1]} yes/no checks on the whole output, and each set "
         f"member {MEMBER_CHECKS_RANGE[0]}–{MEMBER_CHECKS_RANGE[1]} more; every check belongs to one criterion "
         "(Checks tab). " + rubric["scoring"]["criterion"] + " " + rubric["scoring"]["scenario"]),
        ("Pass", rubric["scoring"]["gate"] + " " + rubric["scoring"]["pass"]
         + f" The pass score is {rubric['pass_score']:.1f}. Each family's weights and gate are on the "
         "Rubric by family tab; each scenario's own weights and pass rule are on the Scenario Bank tab."),
        ("Difficulty", "Hard: a strong model gets most of it right and misses a detail an expert would catch. "
         "Extreme: several levers stacked; even the best output is expected to miss something."),
        ("Sources", "Never use a model in the comparison to make its own inputs. Which tool made each source "
         "is recorded in a provenance file when the source is created, not in this bank."),
        ("Coverage", "Rules that need a full family or a full bank show as pending until then; the build "
         "stops on any rule that is already broken."),
    ]
    for head, body in paras:
        ws.append([head, body])
        ws.cell(ws.max_row, 1).font = SECTION_FONT
        ws.cell(ws.max_row, 2).alignment = WRAP
    ws.append([])
    ws.column_dimensions["C"].width = 60
    ws.append(["Hardness lever", "What it means", "Example scenarios"])
    for c in ws[ws.max_row]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    shown = Counter()
    for k, v in LEVERS.items():
        ex = [s for s in scen if k in s["levers"]]
        ex = sorted(ex, key=lambda s: (not s.get("seed"), shown[s["id"]], s["id"]))[:2]
        shown.update(s["id"] for s in ex)
        ws.append([k, v, "\n".join(f"{s['id']}: {s['title']}" for s in ex)])
        ws.cell(ws.max_row, 3).alignment = WRAP
    ws.append([])
    ws.append(["Set mode", "What it means"])
    for c in ws[ws.max_row]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    used = {(s.get("set_members") or {}).get("mode") for s in scen}
    for k, v in SET_MODES.items():
        if k in used:
            ws.append([k, v])

    # 2 · Scenario Bank
    ws = wb.create_sheet("Scenario Bank")
    rows = []
    for s in scen:
        sl, fam = s["_slot"], fams[s["_slot"]["_fam"]]
        sm = s.get("set_members") or {}
        rows.append([
            s["id"], fam["family"], INDUSTRIES[sl["_ind"]], customers_text(customers, s["id"]),
            TASKS[sl["task"]], DIFFS[sl["diff"]], s["title"], inputs_text(s), s["prompt"].strip(),
            SET_MODES.get(sm.get("mode"), "—") if sm else "—", outputs_of(s), members_text(s),
            s["deliverable"], "\n".join(s.get("exact_text") or []) or "—", sum(s["_crits"].values()),
            rubric_text(s, fam, rubric), pass_text(s, fam, rubric), video_text(vlinks, s["id"]),
            SUBJECTS[sl["subject"]], STYLES[sl["style"]], LIGHTS[sl["light"]], LOCALES[sl["locale"]],
            SCRIPTS[sl["script"]],
        ])
    table(ws, ["Scenario ID", "Family", "Industry", "Target customers", "Task", "Difficulty", "Title",
               "Inputs (source id, role)", "Prompt", "Set mode", "Outputs", "What each output is",
               "Deliverable", "Exact text", "Checks", "Rubric (weight · checks)", "Pass rule", "Video follow-on",
               "Subject", "Style", "Light", "Locale", "Script"],
          rows, [13, 20, 20, 30, 12, 10, 30, 30, 90, 18, 8, 70, 34, 30, 8, 40, 50, 50, 14, 14, 14, 12, 12])

    # 2b · Review
    ws = wb.create_sheet("Review")
    rows = [[s["id"], s["title"], s["persona"], target_label(fams[s["_slot"]["_fam"]], s["_slot"]["target"]),
             ", ".join(s["levers"])] for s in scen]
    table(ws, ["Scenario ID", "Title", "Persona", "Failure target", "Levers"], rows, [13, 40, 50, 50, 40])
    customers_sheet(wb, scen, customers, table)

    # 3 · Rubric
    ws = wb.create_sheet("Rubric")
    used = Counter(c[0] for s in scen for c in all_checks(s))
    rows = [[c["name"], k, c["judges"], c["full"], c["half"], c["applies"], used[k]]
            for k, c in rubric["criteria"].items()]
    table(ws, ["Criterion", "Key", "What it judges", "Full marks (10) look like", "Half marks (5) look like",
               "Scored in", "Checks in the bank"], rows, [22, 20, 60, 50, 60, 34, 10])

    # 4 · Rubric by family
    ws = wb.create_sheet("Rubric by family")
    rows = []
    for code, f in fams.items():
        fr = f["rubric"]
        fc = Counter(c[0] for s in f["scenarios"] for c in all_checks(s))
        rows.append([f["family"], "\n".join(f"{rubric['criteria'][k]['name']}: {v}" for k, v in fr["weights"].items()),
                     ", ".join(rubric["criteria"][k]["name"] for k in fr["gate"]), fr["gate_rule"],
                     "\n".join(f"{rubric['criteria'][k]['name']}: {fc[k]}" for k in fr["weights"] if fc[k])])
    table(ws, ["Family", "Weights (points; shares are taken over the criteria a scenario checks)",
               "Gate", "Why it is the gate", "Checks per criterion"], rows, [26, 40, 30, 50, 34])

    # 5 · Checks
    ws = wb.create_sheet("Checks")
    rows = []
    for s in scen:
        gate = fams[s["_slot"]["_fam"]]["rubric"]["gate"]
        for i, (k, q) in enumerate(s["checks"], 1):
            rows.append([s["id"], "Whole output", i, rubric["criteria"][k]["name"], "Yes" if k in gate else "", q])
        for mem in (s.get("set_members") or {}).get("members") or []:
            for i, (k, q) in enumerate(mem.get("checks") or [], 1):
                rows.append([s["id"], mem["name"], i, rubric["criteria"][k]["name"], "Yes" if k in gate else "", q])
    table(ws, ["Scenario ID", "Applies to", "#", "Rubric criterion", "Gate", "Yes/no check"],
          rows, [13, 30, 5, 22, 7, 110])

    # 4 · Set members
    ws = wb.create_sheet("Set members")
    rows = []
    for s in scen:
        sm = s.get("set_members")
        if not sm:
            continue
        rows.append([s["id"], "Must stay the same", "", "\n".join(f"• {x}" for x in sm["must_stay_same"]), ""])
        for j, mem in enumerate(sm["members"], 1):
            rows.append([s["id"], f"{j}. {mem['name']}", SET_MODES[sm["mode"]], mem["prompt"].strip(),
                         "\n".join(f"{rubric['criteria'][d]['name']}: {q}" for d, q in mem["checks"])])
    table(ws, ["Scenario ID", "Member", "Set mode", "Member brief (sent with the shared prompt)", "Member checks"],
          rows, [13, 30, 26, 80, 70])

    # 5 · Slot grid
    ws = wb.create_sheet("Slot grid")
    written = {s["id"] for s in scen}
    rows = []
    for f in fams.values():
        for sl in f["slots"]:
            rows.append([sl["id"], f["family"], INDUSTRIES[sl["_ind"]], TASKS[sl["task"]], DIFFS[sl["diff"]],
                         SUBJECTS[sl["subject"]], STYLES[sl["style"]], LIGHTS[sl["light"]],
                         LOCALES[sl["locale"]], SCRIPTS[sl["script"]], target_label(f, sl["target"]),
                         ", ".join(sl["sources"]) or "—", sl["idea"],
                         "Written" if sl["id"] in written else "Empty"])
    table(ws, ["Slot", "Family", "Industry", "Task", "Difficulty", "Subject", "Style", "Light", "Locale",
               "Script", "Failure target", "Sources", "Idea", "Status"],
          rows, [13, 20, 20, 12, 10, 16, 16, 15, 13, 12, 40, 30, 70, 9])

    # 6 · Scene kits
    ws = wb.create_sheet("Scene kits")
    users = defaultdict(list)
    for sl in slots:
        for sid in sl["sources"]:
            users[sid].append(sl)
    rows = []
    for sid, src in sources.items():
        if src["role"] != "scene":
            continue
        us = users.get(sid, [])
        rows.append([sid, len(us), f"{MAX_KIT_USES - len(us)}", ", ".join(u["id"] for u in us),
                     ", ".join(sorted({fams[u["_fam"]]["family"] for u in us})),
                     "\n".join(v for s in scen if sid in s["_slot"]["sources"]
                               for v, _ in vlinks.get(s["id"], [])) or "—"])
    table(ws, ["Scene kit", "Uses", "Uses left (cap 5, incl. video)", "Used by", "Families", "Video follow-on"],
          rows, [22, 7, 14, 50, 50, 60])

    # 7 · Sources
    ws = wb.create_sheet("Sources")
    rows = []
    for sid, src in sources.items():
        rows.append([sid, src["kind"], src["role"], ", ".join(src.get("used_by") or []), src["aspect"],
                     src["resolution"], src["source_prompt"].strip(),
                     "\n".join(f"• {x}" for x in src.get("avoid") or []),
                     "\n".join(f"☐ {x}" for x in src.get("must_contain") or []), src.get("alt_real") or "—"])
    table(ws, ["Source", "Kind", "Role", "Used by", "Aspect", "Resolution", "Source prompt (model-neutral)",
               "Must not appear", "Acceptance checklist", "Real photo instead"],
          rows, [22, 7, 10, 34, 8, 18, 100, 34, 44, 34])

    # 8 · Task-type split
    ws = wb.create_sheet("Task-type split")
    rows = []
    for code, f in fams.items():
        for ind, (qty, t2i, edit) in PLAN[code].items():
            got = [sl for sl in f["slots"] if sl["_ind"] == ind]
            t = Counter(sl["task"] for sl in got)
            rows.append([f["family"], INDUSTRIES[ind], qty, t2i, edit, len(got), t["t2i"], t["edit"],
                         sum(1 for sl in got if sl["id"] in written)])
    tt = Counter(sl["task"] for sl in slots)
    rows.append(["Total", "", BANK_TOTAL, sum(v[1] for p in PLAN.values() for v in p.values()),
                 sum(v[2] for p in PLAN.values() for v in p.values()), len(slots), tt["t2i"], tt["edit"],
                 len(scen)])
    table(ws, ["Family", "Industry", "Planned", "Planned t2i", "Planned edit", "Slots declared",
               "Declared t2i", "Declared edit", "Written"], rows, [26, 34, 9, 10, 10, 10, 10, 10, 9])

    # 9 · Families
    ws = wb.create_sheet("Families")
    rows = []
    for code, f in fams.items():
        planned = sum(q for q, _, _ in PLAN[code].values())
        sah = [f"{v[0]} ({v[1]})" for k, v in SAHIL.items() if v[2] == code]
        rows.append([f["number"], code, f["family"], f["about"], planned, planned_outputs(code),
                     sum(v[1] for v in PLAN[code].values()), sum(v[2] for v in PLAN[code].values()),
                     "\n".join(f"{k} — {v}" for k, v in (f.get("targets") or {}).items()),
                     "\n".join(sah) or "—",
                     ", ".join(s["id"] for s in f["scenarios"] if s.get("seed")) or "—"])
    table(ws, ["#", "Code", "Family", "What it tests", "Scenarios", "Outputs each", "t2i", "edit",
               "Failure targets", "Sahil's examples (screenshots)", "Seeds written"],
          rows, [4, 7, 24, 50, 9, 9, 6, 6, 70, 40, 30])

    # 10 · Coverage matrix
    ws = wb.create_sheet("Coverage matrix")
    table(ws, ["Scope", "Rule", "Value", "Status"], [list(x) for x in rep.rules], [34, 40, 50, 11])
    for row in ws.iter_rows(min_row=2):
        fill = STATUS_FILL.get(row[3].value)
        if fill:
            row[3].fill = fill

    # 11 · Image → Video links
    ws = wb.create_sheet("Image → Video links")
    rows = [[s["id"], s["title"], ", ".join(s["_slot"]["sources"]) or "—", v, idea]
            for s in scen for v, idea in vlinks.get(s["id"], [])]
    table(ws, ["Image scenario", "Title", "Shared sources", "Video scenario", "What the video does"],
          rows, [13, 40, 34, 13, 80])
    return wb


# ---------------------------------------------------------------- main

def progress_table(fams: OrderedDict) -> str:
    head = "family   " + " ".join(f"{i:>9}" for i in INDUSTRIES) + "     total"
    lines = [head]
    for code, f in fams.items():
        cells = []
        for ind in INDUSTRIES:
            if ind not in PLAN[code]:
                cells.append(f"{'·':>9}")
                continue
            n_slots = sum(1 for sl in f["slots"] if sl["_ind"] == ind)
            n_written = sum(1 for s in f["scenarios"] if s.get("_slot", {}).get("_ind") == ind)
            cells.append(f"{n_written:>2}/{n_slots:>2}/{PLAN[code][ind][0]:>2}")
        tot = sum(q for q, _, _ in PLAN[code].values())
        lines.append(f"{code:<8} " + " ".join(cells) + f"   {len(f['scenarios']):>3}/{len(f['slots']):>3}/{tot}")
    return "\n".join(lines) + "\n(written / slots declared / planned)"


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--src", type=Path, default=SRC)
    ap.add_argument("--out", type=Path, default=OUT)
    ap.add_argument("--check", action="store_true", help="validate only; do not write the workbook")
    a = ap.parse_args()

    fams, sources = load(a.src)
    rubric = load_rubric(a.src / "rubric.yaml")
    customers = load_customers(a.src / "customers.yaml")
    rep = validate(fams, sources, rubric, customers)
    if rep.errors:
        print(f"{len(rep.errors)} rule(s) broken:", file=sys.stderr)
        for e in rep.errors:
            print(f"  - {e}", file=sys.stderr)
        return 1
    print(progress_table(fams))
    pend = [x for x in rep.rules if x[3] in ("pending", "behind")]
    print(f"\n{len(rep.rules) - len(pend)} coverage rules met or on track, {len(pend)} pending.")
    scen = [s for f in fams.values() for s in f["scenarios"]]
    print(f"{len(scen)} scenarios, {sum(outputs_of(s) for s in scen)} outputs per model, {len(sources)} sources.")
    if a.check:
        return 0
    wb = build(fams, sources, rubric, rep, customers)
    a.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(a.out)
    print(f"wrote {a.out.relative_to(ROOT) if a.out.is_relative_to(ROOT) else a.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
